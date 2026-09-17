"""Quantitatif traçable : lignes métier reliées aux preuves et calcul décimal.

Aucune quantité issue des pixels n'est promue en mesure exacte. Les postes sans
cotes ou sans correspondance CCTP sont rendus dans les réserves, pas inventés.
"""
import asyncio,hashlib,json,re
from decimal import Decimal
from ressources import dossiers
from skills.registre import Declaration
from skills.documents_dossier import _json,_analyses
from skills.chiffres_sources import calculer,NOMBRE,VALEUR


def nombre(texte):
    return Decimal(str(texte).replace(' ','').replace('\u00a0','').replace('\u202f','').replace(',','.'))

def _verifier_ligne(ligne,preuves):
    for k in ('lot','poste','niveau','local','unite'):
        if not isinstance(ligne.get(k),str) or not ligne[k].strip():raise ValueError('Ligne quantitative sans '+k)
    if ligne['unite'] not in ('m²','m','ml','m³','u','kg'):raise ValueError('Unité de résultat inconnue.')
    valeurs={};references=[]
    operandes=ligne.get('operandes')
    if not isinstance(operandes,list) or not 1<=len(operandes)<=12:raise ValueError('De 1 à 12 opérandes sourcés sont nécessaires.')
    for i,o in enumerate(operandes):
        extrait=preuves.get(o.get('preuve'))
        citation=str(o.get('citation') or '')
        if not extrait or not citation or ' '.join(citation.split()) not in ' '.join(extrait.split()):raise ValueError('Citation de quantité absente de la source.')
        valeur=nombre(o.get('valeur'))
        trouves=[]
        for m in NOMBRE.finditer(citation):
            v=VALEUR.match(m.group());trouves.append(nombre(v[1]))
        if valeur not in trouves:raise ValueError('La valeur ne figure pas dans la citation donnée.')
        unite=o.get('unite')
        if unite not in ('m²','m','ml','m³','u','kg'):raise ValueError('Unité d’opérande inconnue ; ne déduis pas une unité absente.')
        # L'unité est citée dans le même extrait (en-tête de tableau compris).
        alias={'m²':('m²','m2'),'m³':('m³','m3'),'ml':('ml','m'),'m':('m',),'u':('unité','unite',' u','quantité','quantite'),'kg':('kg',)}[unite]
        if not any(re.search(r'(?<![\w])'+re.escape(a.strip())+r'(?![\w])',citation.lower()) for a in alias):raise ValueError('L’unité doit apparaître dans la citation, avec la valeur.')
        valeurs['c'+str(i+1)]=str(valeur);references.append(o)
    resultat=Decimal(calculer(ligne.get('formule') or 'c1',valeurs))
    if resultat<0:raise ValueError('Une quantité négative nécessite de revoir les déductions.')
    justification=ligne.get('affectation') or {}
    if not justification.get('preuve') in preuves or not justification.get('citation'):raise ValueError('L’affectation au poste nécessite une preuve du CCTP, DPGF ou tableau de localisation.')
    if ' '.join(justification['citation'].split()) not in ' '.join(preuves[justification['preuve']].split()):raise ValueError('Affectation non prouvée par les pièces.')
    # Vérification dimensionnelle : surfaces additionnées ; longueurs multipliées.
    import ast
    dimensions={'m':(1,0),'ml':(1,0),'m²':(2,0),'m³':(3,0),'u':(0,0),'kg':(0,1)}
    dims={'c'+str(i+1):dimensions[o['unite']] for i,o in enumerate(operandes)}
    def dimension(n):
        if isinstance(n,ast.Name):return dims[n.id]
        if isinstance(n,ast.UnaryOp):return dimension(n.operand)
        if isinstance(n,ast.BinOp):
            a,b=dimension(n.left),dimension(n.right)
            if isinstance(n.op,(ast.Add,ast.Sub)):
                if a!=b:raise ValueError('Addition de grandeurs incompatibles.')
                return a
            if isinstance(n.op,ast.Mult):return tuple(x+y for x,y in zip(a,b))
            if isinstance(n.op,ast.Div):return tuple(x-y for x,y in zip(a,b))
        raise ValueError('Formule dimensionnelle invalide.')
    if dimension(ast.parse(ligne.get('formule') or 'c1',mode='eval').body)!=dimensions[ligne['unite']]:raise ValueError('Unité du résultat incompatible avec le calcul.')
    return {**ligne,'quantite':format(resultat,'f')}

def dedoublonner(lignes):
    retenues={};conflits=[];bloquees=set()
    for l in lignes:
        cle=tuple(' '.join(l[k].casefold().split()) for k in ('lot','poste','niveau','local','unite'))
        if cle in bloquees:continue
        if cle in retenues and retenues[cle]['quantite']!=l['quantite']:
            conflits.append('Quantités contradictoires pour '+' / '.join(cle));retenues.pop(cle);bloquees.add(cle)
        else:retenues.setdefault(cle,l)
    return list(retenues.values()),conflits

async def _produire(data,user):
    uid,fil=dossiers.identite(getattr(user,'id',None),data.get('_fil'))
    demande=str(data.get('_demande_utilisateur') or data.get('demande') or '')
    ids=data.get('sources') or [s['id'] for s in dossiers.manifeste(uid,fil)]
    if not ids:raise ValueError('Ajoute les plans, CCTP et tableaux au dossier avant de demander le quantitatif.')
    sources=await asyncio.to_thread(dossiers.sources,uid,fil,ids)
    from skills.documents_dossier import verifier_acces
    await verifier_acces(sources,user)
    tache=data['_tache_quantitatif']
    from skills.documents_dossier import completer_visuels
    sources=await completer_visuels(uid,fil,tache,sources,user,demande)
    analyses=await _analyses(uid,fil,tache,demande,sources)
    # Les contraintes et affectations sont communes à tous les fragments à quantifier.
    contexte=[a for a in analyses if any(f.get('nature')=='exigence' for f in a['faits'])]
    fragments=[];preuves={}
    for s in sources:
        for f in dossiers.fragments(s['contenu'],taille=5000):
            cle=s['id']+':q'+str(f['numero']);texte=f['texte']
            fragments.append((s,cle,texte));preuves[cle]=texte
        for f in dossiers.fragments(s['contenu']):preuves[f"{s['id']}:{f['numero']}"]=f['texte']
    semaphore=asyncio.Semaphore(3)
    async def extraire(s,cle,texte):
        ancienne=await asyncio.to_thread(dossiers.etape,uid,fil,tache,cle)
        if ancienne:return ancienne
        async with semaphore:
            def verifier(r):
                if not isinstance(r.get('lignes'),list) or not isinstance(r.get('reserves'),list):raise ValueError('lignes[] et reserves[] obligatoires.')
                r['lignes']=[_verifier_ligne(l,preuves) for l in r['lignes']]
            r=await _json('Établis les lignes de quantitatif prouvables à partir de CE fragment et des affectations du dossier. '
                'Ne réemploie pas les quantités d’un ancien exemple. Ne confonds pas surface habitable et surface de revêtement. '
                'Pas de faïence calculée sans hauteur/périmètre, pas de plinthes sans longueur justifiée. '
                'Un même local/poste doit avoir une clé stable entre plans et tableau. Donne toutes les lignes pertinentes du fragment ; '
                'réserves précises pour les postes non déterminables. Schéma {"lignes":[{"lot":"...","poste":"...","niveau":"...",'
                '"local":"...","unite":"m²","formule":"c1","operandes":[{"valeur":"42.5","unite":"m²","preuve":"référence",'
                '"citation":"texte exact portant valeur et unité"}],"affectation":{"preuve":"référence","citation":"texte exact du CCTP/localisation"}}],"reserves":["..."]}. '
                'La formule emploie UNIQUEMENT c1, c2, etc. dans l’ordre des opérandes et les signes + - * / : '
                'pour longueur fois largeur, écris c1*c2, jamais longueur x largeur. '
                'Les identifiants, dates et numéros de plan ne sont pas des quantités. Une estimation visuelle est une réserve, pas une mesure exacte.',
                {'demande':demande,'source':s['nom'],'preuve_fragment':cle,'fragment':texte,'affectations':contexte},verifier)
            await asyncio.to_thread(dossiers.etape,uid,fil,tache,cle,r);return r
    lus=await asyncio.gather(*(extraire(s,c,t) for s,c,t in fragments),return_exceptions=True)
    erreurs=['Un fragment n’a pas pu être quantifié ('+type(x).__name__+').' for x in lus if isinstance(x,BaseException)]
    bons=[x for x in lus if isinstance(x,dict)]
    lignes,conflits=dedoublonner([l for r in bons for l in r['lignes']])
    reserves=list(dict.fromkeys([x for r in bons for x in r['reserves']]+conflits+erreurs))
    visuels={s['id'] for s in sources if 'lecture visuelle' in s['nom'].casefold() or 'analyse visuelle' in s['contenu'][:150].casefold()}
    for l in lignes:
        l['lecture']='Visuelle : à contrôler sur le plan' if any(o['preuve'].split(':')[0] in visuels for o in l['operandes']) else 'Texte extrait'
    if any(l['lecture'].startswith('Visuelle') for l in lignes):reserves.append('Les quantités issues de lectures visuelles doivent être contrôlées sur les plans originaux ; la citation vérifie la transcription conservée, pas la mesure physique.')
    if lignes:
        avis=await _json('Vérifie chaque ligne quantitative : affectation réelle au poste et au local, absence de double comptage, ancienne opération non réemployée, cotes et unités lisibles. '
            'Pour chaque ligne invalide, donne son indice zéro-based dans rejeter et une réserve précise. JSON {"rejeter":[0],"reserves":["..."]}. Une simple surface habitable ne prouve pas automatiquement une surface de revêtement.',
            {'demande':demande,'lignes':lignes,'affectations':contexte})
        rejeter=avis.get('rejeter');notes=avis.get('reserves')
        if not isinstance(rejeter,list) or any(type(i) is not int or not 0<=i<len(lignes) for i in rejeter) or not isinstance(notes,list):raise ValueError('Contrôle quantitatif invalide.')
        lignes=[l for i,l in enumerate(lignes) if i not in rejeter];reserves.extend(str(n) for n in notes)
    if not lignes:return {'ok':True,'outcome':'partial','production_verifiee':False,'reserves':reserves or ['Aucune quantité avec affectation et unité suffisamment prouvées.'],'a_faire':'Identifie les pages et cotes manquantes ; analyser_plan_source permet de lire le dessin. Ne livre pas un inventaire de fichiers comme métré.'}
    from bureautique import atelier
    from skills.bureau import terminer_document
    titres=['Lot','Poste','Niveau','Local','Quantité','Unité','Formule','Sources et citations','Lecture']
    rows=[[l[k] for k in ('lot','poste','niveau','local','quantite','unite','formule')]+[json.dumps({'operandes':l['operandes'],'affectation':l['affectation']},ensure_ascii=False),l['lecture']] for l in lignes]
    for row in rows:row[4]=float(Decimal(row[4]))
    totaux={}
    for l in lignes:
        k=(l['lot'],l['poste'],l['niveau'],l['unite']);totaux[k]=totaux.get(k,Decimal(0))+Decimal(l['quantite'])
    blocs=[{'bloc':'feuille','nom':'Détail','colonnes_numeriques':[4],'entetes':titres,'lignes':rows},
           {'bloc':'feuille','nom':'Synthèse','colonnes_numeriques':[4],'entetes':['Lot','Poste','Niveau','Unité','Quantité'],'lignes':[[*k,float(v)] for k,v in sorted(totaux.items())]},
           {'bloc':'feuille','nom':'Réserves','entetes':['Point à vérifier'],'lignes':[[r] for r in reserves] or [['Aucune réserve détectée ; contrôle métier humain requis.']]}]
    preuves_lignes=[]
    for i,l in enumerate(lignes,1):
        for o in l['operandes']:
            for debut in range(0,len(o['citation']),1900):preuves_lignes.append([i,'Opérande',o['preuve'],o['valeur'],o['unite'],o['citation'][debut:debut+1900]])
        a=l['affectation']
        for debut in range(0,len(a['citation']),1900):preuves_lignes.append([i,'Affectation',a['preuve'],'','',a['citation'][debut:debut+1900]])
    for debut in range(0,len(preuves_lignes),5000):
        blocs.append({'bloc':'feuille','nom':'Preuves '+str(debut//5000+1),'entetes':['Ligne détail','Type','Source et fragment','Valeur','Unité','Citation'],'lignes':preuves_lignes[debut:debut+5000]})
    jeton=await asyncio.to_thread(dossiers.etape,uid,fil,tache,'jeton')
    if not jeton:
        jeton=await asyncio.to_thread(atelier.ouvrir,{'titre':data.get('titre') or 'Quantitatif sourcé','format':'xlsx'},uid,fil)
        await asyncio.to_thread(dossiers.etape,uid,fil,tache,'jeton',jeton)
    from bureautique.modele import normaliser_element
    attendus=[normaliser_element(b) for b in blocs];presents=list(atelier.elements(jeton))
    if presents!=attendus[:len(presents)]:
        await asyncio.to_thread(atelier.abandonner,jeton,uid)
        await asyncio.to_thread(dossiers.effacer_etapes,uid,fil,tache,['jeton'])
        raise ValueError('Le quantitatif a évolué ; nouveau rendu à la prochaine reprise.')
    if len(presents)<len(attendus):await asyncio.to_thread(atelier.ajouter,jeton,attendus[len(presents):],uid,False)
    r=await terminer_document({'document_id':jeton,'_fil':fil},user)
    from openpyxl import load_workbook
    wb=load_workbook(atelier.chemin_fichier(jeton,uid),read_only=True,data_only=True)
    try:
        if wb['Détail'].max_row!=len(lignes)+1:raise ValueError('Le classeur rendu ne contient pas toutes les lignes calculées.')
        if any(type(row[4]) not in (int,float) for row in wb['Détail'].iter_rows(min_row=2,values_only=True)):raise ValueError('Les quantités Excel ne sont pas numériques.')
    finally:wb.close()
    return {**r,'ok':True,'production_verifiee':True,'outcome':'partial' if reserves else 'success','lignes':len(lignes),
            'reserves':reserves,'sources_lues':len(sources),'fragments_quantifies':len(bons),'fragments_total':len(fragments),'a_faire':'Présente le quantitatif, sa couverture et les réserves. Il contient les quantités établies avec leurs preuves et calculs ; ne le prétends pas exhaustif si des postes restent en réserve.'}

async def produire_immediat(data,user):
    uid,fil=dossiers.identite(getattr(user,'id',None),data.get('_fil'))
    data=await asyncio.to_thread(dossiers.normaliser_selection,uid,fil,data)
    ids=data.get('sources') or [s['id'] for s in dossiers.manifeste(uid,fil)]
    demande=str(data.get('_demande_utilisateur') or data.get('demande') or '')
    tache=hashlib.sha256(json.dumps(['quantitatif',demande,ids,data.get('titre')],ensure_ascii=False).encode()).hexdigest()[:24]
    from ressources.documents_file import associer_tache
    await asyncio.to_thread(associer_tache,uid,fil,tache)
    from stockage.verrous import verrou_fichier
    from ressources.registre import _chemin
    from bureautique import atelier
    with verrou_fichier(str(_chemin().parent),'document:'+uid+':'+tache,bloquant=False) as acquis:
        if not acquis:return {'ok':True,'en_cours':True,'production_verifiee':False,'tache':tache}
        fini=await asyncio.to_thread(dossiers.etape,uid,fil,tache,'livraison')
        if fini:
            if not atelier.chemin_fichier(fini['document_id'],uid):raise ValueError('Le quantitatif sauvegardé est devenu indisponible.')
            return fini
        await asyncio.to_thread(dossiers.etape,uid,fil,tache,'contrat',{'demande':demande,'sources':ids})
        try:r=await _produire({**data,'sources':ids,'_tache_quantitatif':tache},user)
        except Exception as e:
            # La file doit mesurer les étapes de ce quantitatif même après échec.
            e.tache_documentaire=tache
            raise
        r['tache']=tache
        if r.get('production_verifiee'):await asyncio.to_thread(dossiers.etape,uid,fil,tache,'livraison',r)
        return r

async def produire(data,user):
    from ressources.documents_file import soumettre
    uid,fil=dossiers.identite(getattr(user,'id',None),data.get('_fil'))
    return await asyncio.to_thread(soumettre,uid,fil,'quantitatif',data)

SKILLS={'produire_quantitatif' :Declaration(produire,'Produire un Excel de quantités depuis les plans, CCTP/DPGF et tableaux du dossier : détail par poste/local/niveau, unités et calculs vérifiés, synthèse et réserves. Aucun chiffre inventé. Ajouter les pièces avec ajouter_source_dossier et lire les plans graphiques avec analyser_plan_source avant si nécessaire.',requis=['demande'],optionnels=['sources','titre'],effet='ecriture_interne',expert='agent2',libelle='je calcule le quantitatif et vérifie ses sources')}
