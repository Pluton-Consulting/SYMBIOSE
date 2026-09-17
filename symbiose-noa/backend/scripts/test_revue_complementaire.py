"""Régressions de la revue du 16/09 : vrais DOCX et contrats exécutés hors réseau.

Les cas portent sur des effets observables, pas sur la présence d'une ligne
dans le source. Les services externes sont remplacés ; aucun identifiant réel.
"""
import ast
import asyncio
import importlib.util
import io
import json
import logging
import os
from pathlib import Path
import random
import sys
import types
import zipfile
from unittest.mock import patch

B = Path(sys.argv[1] if len(sys.argv) > 1 else "backend").resolve()
sys.path.insert(0, str(B))
ECHECS = []


def verifier(nom, fonction):
    try:
        fonction()
        print("✓", nom)
    except Exception as e:
        ECHECS.append(nom)
        print("✗", nom, type(e).__name__, str(e)[:300])


def charger(nom, chemin):
    spec = importlib.util.spec_from_file_location(nom, B / chemin)
    mod = importlib.util.module_from_spec(spec)
    sys.modules[nom] = mod
    spec.loader.exec_module(mod)
    return mod


def extraire(chemin, noms, env):
    arbre = ast.parse((B / chemin).read_text())
    nodes = [n for n in arbre.body if isinstance(n, (ast.FunctionDef, ast.AsyncFunctionDef)) and n.name in noms]
    module = ast.Module(body=[ast.ImportFrom(module="__future__", names=[ast.alias(name="annotations")], level=0)] + nodes, type_ignores=[])
    exec(compile(ast.fix_missing_locations(module), str(chemin), "exec"), env)


def documents():
    from docx import Document
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    tr = charger("trame_revue", "bureautique/trame.py")
    co = charger("controle_revue", "bureautique/controle.py")
    def sauver(d):
        b = io.BytesIO(); d.save(b); return b.getvalue()
    d = Document(); p = d.add_paragraph("Client MARTIN ")
    champ = OxmlElement("w:fldSimple"); champ.set(qn("w:instr"), "DATE")
    run = OxmlElement("w:r"); texte = OxmlElement("w:t"); texte.text = "15/09/2026"
    run.append(texte); champ.append(run); p._p.append(champ)
    d.add_paragraph()
    b = io.BytesIO()
    with zipfile.ZipFile(io.BytesIO(sauver(d))) as zi, zipfile.ZipFile(b, "w") as zo:
        for f in zi.infolist():
            contenu = zi.read(f.filename)
            if f.filename == "word/document.xml":
                contenu = contenu.replace(b"<w:p/>", b"<w:p></w:p>")
            zo.writestr(f, contenu)
    out = tr.remplir_detaille(b.getvalue(), "docx", {"MARTIN": "DUPONT\n12 rue A\tParis", "15/09/2026": "20/09/2026"})
    with zipfile.ZipFile(io.BytesIO(out["octets"])) as z:
        contenu = z.read("word/document.xml")
    assert b"15/09/2026" in contenu and b"20/09/2026" not in contenu
    assert contenu.count(b"<w:br") == 1 and contenu.count(b"<w:tab") == 1
    assert out["remplacements"] == 1 and out["controle"]["ok"]
    d = Document(); d.add_paragraph("Texte contractuel")
    avant = sauver(d); d.paragraphs[0].runs[0].text = "Autre texte"; d.sections[0].left_margin = 1
    assert not co.comparer_docx(avant, sauver(d))["ok"]
    rng = random.Random(42)
    for _ in range(120):
        text = " ".join(rng.choice(["MARTIN", "DUPONT", "AUTRE"]) for _ in range(7))
        d = Document(); p = d.add_paragraph(); pos = 0
        while pos < len(text):
            taille = rng.randint(1, 9); p.add_run(text[pos:pos + taille]); pos += taille
        tr._remplacer_dans_paragraphe(p, {"MARTIN": "A\nB", "DUPONT": "C\tD"})
        assert p.text == text.replace("MARTIN", "A\nB").replace("DUPONT", "C\tD")


def statuts():
    r = charger("skills.resultats", "skills/resultats.py")
    assert r.normaliser_resultat({"message": "libre"}, "ecriture_interne")["effect_status"] == "unknown"
    for statut in ("pending", "not_found", "failed", "denied", "unverified"):
        assert not r.resultat_probant({"ok": True, "outcome": statut})
    assert r.resultat_probant({"ok": True, "outcome": "success"})
    env = {}
    extraire("agents/agent1.py", {"resultat_de_geste"}, env)
    agent = types.ModuleType("agents.agent1")
    agent.resultat_de_geste = env["resultat_de_geste"]
    agent.PLAFOND_RESULTAT = 4000; agent.PLAFOND_RESULTAT_GENEREUX = 12000; agent.RESULTATS_GENEREUX = set()
    ex = types.ModuleType("skills.executor"); ex.expert_du_skill = lambda s: None
    an = types.ModuleType("security.anonymizer")
    an.anonymizer = types.SimpleNamespace(anonymize_chunks=lambda chunks, carte: (chunks, carte))
    with patch.dict(sys.modules, {"skills.executor": ex, "agents.agent1": agent, "security.anonymizer": an}):
        env = {}; extraire("agents/router.py", {"_reprise_du_tour", "_reouverture_du_tour"}, env)
        envelope = {"output": {"statut": "en_attente"}, **r.normaliser_resultat({"statut": "en_attente"}, "externe")}
        out = asyncio.run(env["_reprise_du_tour"]({}, {"skill": "deposer_brouillon"}, "x", envelope, None))
        entree = out["tool_results"][0]
        assert entree["outcome"] == "pending" and entree["effect_status"] == "pending"
        assert not r.resultat_probant(entree)


def recherche():
    import vectorstore.fusion as fusion
    faux = types.ModuleType("vectorstore.rag")
    auth = types.ModuleType("mail.authorization")
    async def boites(_): return []
    auth.boites_pour_la_memoire = boites
    auth.boites_par_id = boites
    async def inventaire(_): return "12 documents", False
    env = {"MAX_RESULTATS": 6, "MAX_LIMITE": 20, "BUDGET_EXTRAITS": 9000,
           "logger": logging.getLogger("revue"), "_inventaire": inventaire}
    extraire("skills/documents.py", {"_entier", "rechercher_documents"}, env)
    async def vide(*a, **k):
        return {"documents": [], "diagnostic": {"embedding": "indisponible", "voies": ["texte"]}}
    faux.rechercher = vide
    with patch.dict(sys.modules, {"vectorstore.rag": faux, "vectorstore.fusion": fusion, "mail.authorization": auth}):
        user = types.SimpleNamespace(id="test", role="terrain")
        out = asyncio.run(env["rechercher_documents"]({"requete": "drainage"}, user))
        assert not out["couverture"]["recherche_par_le_sens"]
        assert "sens" in out["message"] and out["diagnostic"]["embedding"] == "indisponible"
        async def panne(*a, **k): raise RuntimeError("panne simulée")
        faux.rechercher = panne
        assert asyncio.run(env["rechercher_documents"]({"requete": "x"}, user))["ok"] is False


def isolement():
    config = types.ModuleType("config"); config.settings = types.SimpleNamespace(daytona_api_key="")
    with patch.dict(sys.modules, {"config": config}), patch.dict(os.environ, {"AUTORISER_CODE_NON_ISOLE": ""}):
        mod = charger("sandbox_revue", "sandbox/daytona_client.py")
        client = mod.DaytonaClient()
        async def interdit(*a, **k): raise AssertionError("Sous-processus local interdit")
        with patch.object(asyncio, "create_subprocess_exec", interdit):
            assert not asyncio.run(client.test_skill("raise Exception()", "test")).passed
            assert not asyncio.run(client.execute_skill("raise Exception()", "test", {}))["ok"]
            client.daytona_available = True
            appels = []
            async def distant(code, timeout):
                appels.append(code)
                return types.SimpleNamespace(exit_code=0, result='__SKILL_RESULT__{"total": 42}')
            client._executer_isole = distant
            r = asyncio.run(client.execute_skill("def run(data): return data", "test", {"total": 42}))
            assert appels and r["sandbox_type"] == "daytona" and r["output"] == {"total": 42}


def secrets():
    m = charger("secrets_revue", "security/secrets.py")
    valeur = "1//jeton-fictif-a-ne-pas-journaliser"
    r = m.masquer_arbre({"refresh_token": valeur, "niveau": {"authorization": "Bearer " + valeur}})
    assert valeur not in json.dumps(r)
    assert valeur not in m.masquer("refresh_token=" + valeur)
    assert "abc.def.ghi" not in m.masquer("Authorization: Bearer abc.def.ghi")
    coffre = charger("coffre_revue", "security/coffre.py")
    with patch.object(coffre, "_cles", return_value=[]):
        try: coffre.chiffrer("secret-fictif", "test")
        except RuntimeError: pass
        else: raise AssertionError("Secret écrit en clair")


def reprise():
    m = charger("requetes_revue", "agents/requetes.py")
    r = m.reponse_de_reprise({"etat": "terminee", "thread_id": "fil", "resultat": json.dumps({"response": "Fait", "status": "ok"})})
    assert r["response"] == "Fait" and r["reprise"] is False
    r = m.reponse_de_reprise({"etat": "en_cours", "thread_id": "fil"})
    assert r["reprise"] is True
    r = m.reponse_de_reprise({"etat": "terminee", "resultat": json.dumps({"status": "pending_validation", "validation_id": "v"})})
    assert r["validation_id"] == "v" and r["status"] == "pending_validation"


for nom, test in [("Word : champs, XML équivalent, sauts, contrôle et 120 fragments", documents),
                  ("Statuts conservés après accord et sans fausse preuve", statuts),
                  ("Recherche vide dégradée et panne explicite", recherche),
                  ("Exécution ET test du code généré hors du backend", isolement),
                  ("Secrets structurés et refus du clair", secrets),
                  ("Reprise du résultat final ou de l'accord", reprise)]:
    verifier(nom, test)
print(f"{'✗' if ECHECS else '✓'} {len(ECHECS)} échec(s)")
sys.exit(bool(ECHECS))
