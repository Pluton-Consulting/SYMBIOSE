"""
Banc du 21/09 — le Word du dossier Camp et « les 10 autres photos ».

1. LE WORD. « Refais le dossier de M. Camp en Word avec l'en-tête du PDF
   symbiose_devisfinal » : 22 sections versées, puis trois échecs d'affilée,
   « Le document n'a pas pu être produit () ». Le logo tiré du PDF est un JPEG RGB
   qui s'ouvre sur un marqueur Adobe (FFD8 FFEE) ; python-docx ne reconnaît un JPEG
   qu'à FFD8 FFE0 (JFIF) ou FFD8 FFE1 (Exif) et lève `UnrecognizedImageError`, une
   exception SANS message. Le rangement convertit désormais ce que Word ne lit pas,
   le rendu convertit ce qui avait été rangé avant, et l'échec dit sa cause.

2. LES PHOTOS. « Montre les 10 autres photos du dossier » (22 photos) : le geste
   rendait toujours les 12 plus récentes, sans page suivante. `page` donne la suite.

Les rendus sont EXÉCUTÉS avec python-docx et Pillow quand ils sont là (image du
backend) ; sinon cette partie est sautée et le banc le dit.

Usage : python backend/scripts/test_word_et_photos.py [backend]
"""
import asyncio
import importlib.util
import io
import os
import pathlib
import sys
import tempfile
import types

BACKEND = pathlib.Path(sys.argv[1] if len(sys.argv) > 1 else "backend").resolve()
sys.path.insert(0, str(BACKEND))
echecs = []


def verifier(nom, cond, detail=""):
    print(f"  {'✓' if cond else '✗'} {nom}" + (f"  → {detail}" if detail and not cond else ""))
    if not cond:
        echecs.append(nom)


def charger(nom, chemin):
    spec = importlib.util.spec_from_file_location(nom, chemin)
    module = importlib.util.module_from_spec(spec)
    sys.modules[nom] = module
    spec.loader.exec_module(module)
    return module


print(f"\n═══ LE WORD DU DOSSIER CAMP, LES PHOTOS SUIVANTES — {BACKEND.parent}\n")

# ── 1. Ce que Word reconnaît ─────────────────────────────────────────────────
images = charger("bureautique.images", BACKEND / "bureautique" / "images.py")
verifier("un JPEG JFIF est lisible par Word", images.lisible_par_word(b"\xff\xd8\xff\xe0\x00\x10JFIF"))
verifier("un JPEG Exif est lisible par Word", images.lisible_par_word(b"\xff\xd8\xff\xe1\x00\x10Exif"))
verifier("un PNG est lisible par Word", images.lisible_par_word(b"\x89PNG\r\n\x1a\n...."))
verifier("le logo du PDF (JPEG ouvert sur un marqueur Adobe) ne l'est PAS",
         not images.lisible_par_word(b"\xff\xd8\xff\xee\x00\x0eAdobe"))
verifier("un JPEG ouvert sur sa table de quantification ne l'est pas",
         not images.lisible_par_word(b"\xff\xd8\xff\xdb\x00\x84"))

bureau = (BACKEND / "skills" / "bureau.py").read_text(encoding="utf-8")
verifier("un échec de rendu sans message dit au moins le nom de l'exception",
         "cause = str(e) or type(e).__name__" in bureau and "exc_info=True" in bureau)
verifier("… et nomme une image quand c'est une image",
         "n'est pas lisible par Word" in bureau)

try:
    from PIL import Image
    import docx
    from docx.shared import Cm
except ImportError:
    Image = None
    print("  · Pillow / python-docx absents ici : les rendus ne sont pas exécutés (ils le sont dans l'image)")

if Image is not None:
    def jpeg_sans_jfif(taille=(599, 397)) -> bytes:
        """Un JPEG valide dont le segment JFIF est retiré : il s'ouvre sur FFD8 FFDB."""
        tampon = io.BytesIO()
        Image.new("RGB", taille, (30, 120, 60)).save(tampon, "JPEG", quality=90)
        b = tampon.getvalue()
        assert b[2:4] == b"\xff\xe0"
        longueur = int.from_bytes(b[4:6], "big")
        return b[:2] + b[4 + longueur:]

    brut = jpeg_sans_jfif()
    verifier("le JPEG d'essai est valide pour Pillow", Image.open(io.BytesIO(brut)).size == (599, 397))
    doc = docx.Document()
    try:
        doc.add_paragraph().add_run().add_picture(io.BytesIO(brut), height=Cm(1.2))
        refuse = False
    except Exception as e:  # noqa: BLE001
        refuse = type(e).__name__ == "UnrecognizedImageError" and str(e) == ""
    verifier("python-docx le REFUSE, sans message — le défaut du 21/09 reproduit", refuse)

    source_images = (BACKEND / "bureautique" / "images.py").read_text(encoding="utf-8")
    verifier("le rangement ne garde un JPEG tel quel que s'il est lisible par Word",
             "and lisible_par_word(octets))" in source_images)

    # Le rendu : une image RANGÉE AVANT le correctif est convertie au moment du Word.
    with tempfile.TemporaryDirectory() as dossier:
        atelier = types.ModuleType("bureautique.atelier")
        atelier.chemin_image = lambda nom: os.path.join(dossier, nom) if os.path.exists(os.path.join(dossier, nom)) else None
        sys.modules["bureautique.atelier"] = atelier
        rendu = charger("bureautique.rendu", BACKEND / "bureautique" / "rendu.py")
        nom = "rzhFZWHNessai.img1.jpg"
        with open(os.path.join(dossier, nom), "wb") as f:
            f.write(brut)
        chemin = rendu._image(nom)
        verifier("le rendu donne à Word une copie PNG de l'image illisible",
                 chemin and chemin.endswith(".img1.jpg.word.png") and os.path.exists(chemin), str(chemin))
        verifier("… qui commence par « <jeton>.img » (elle part avec le document)",
                 os.path.basename(chemin or "").startswith("rzhFZWHNessai.img"))
        doc = docx.Document()
        try:
            doc.add_paragraph().add_run().add_picture(chemin, height=Cm(1.2))
            passe = True
        except Exception as e:  # noqa: BLE001
            passe = False
        verifier("… et Word l'accepte", passe)
        with open(os.path.join(dossier, "bon.img2.jpg"), "wb") as f:
            tampon = io.BytesIO()
            Image.new("RGB", (40, 40), (200, 30, 30)).save(tampon, "JPEG")
            f.write(tampon.getvalue())
        verifier("une image déjà lisible est rendue telle quelle, sans copie",
                 rendu._image("bon.img2.jpg") == os.path.join(dossier, "bon.img2.jpg"))

# ── 2. Les photos, page par page ─────────────────────────────────────────────
for nom in ("visuels", "visuels.depot"):
    sys.modules.setdefault(nom, types.ModuleType(nom))
depot = sys.modules["visuels.depot"]
depot.deposer_octets = lambda octets, mime: "cle-" + octets.decode()
drive = charger("outils.drive_banc", BACKEND / "outils" / "drive.py")

FICHIERS = [{"id": f"p{i:02d}", "name": f"IMG_{1673 - i}.JPG", "mimeType": "image/jpeg",
             "size": "1000", "modifiedTime": f"2026-09-09T12:{59 - i:02d}:00Z"} for i in range(22)]


class Requete:
    def __init__(self, fn):
        self.fn = fn

    def execute(self):
        return self.fn()


class Fichiers:
    def list(self, **kw):
        # Le Drive rend 200 par page au plus : on sert en deux pages pour éprouver la suite.
        debut = int(kw.get("pageToken") or 0)
        fin = debut + 15
        return Requete(lambda: {"files": FICHIERS[debut:fin],
                                **({"nextPageToken": str(fin)} if fin < len(FICHIERS) else {})})

    def get_media(self, fileId):
        return Requete(lambda: fileId.encode())


class Service:
    def files(self):
        return Fichiers()


async def _service(_identite=None):
    return Service()

drive._service = _service
perimetres = [("racine", "Drive partagé")]


async def jouer():
    p1 = await drive.photos(None, None, 12, perimetres=perimetres, page=1)
    p2 = await drive.photos(None, None, 12, perimetres=perimetres, page=2)
    return p1, p2

p1, p2 = asyncio.run(jouer())
verifier("le dossier est listé EN ENTIER (22, sur deux pages du Drive)", p1["disponibles"] == 22, str(p1.get("disponibles")))
verifier("page 1 : les 12 plus récentes", [i["legende"] for i in p1["photos"]][:1] == ["IMG_1673.JPG"] and p1["nombre"] == 12)
verifier("page 1 dit la suite (page 2, 10 restantes)", "page=2" in (p1.get("pour_continuer") or "")
         and "10 restante" in p1["pour_continuer"], str(p1.get("pour_continuer")))
verifier("page 2 : les 10 AUTRES, aucune déjà montrée",
         p2["nombre"] == 10 and not {i["cle"] for i in p1["photos"]} & {i["cle"] for i in p2["photos"]})
verifier("page 2 se situe (« 13 à 22 sur 22 ») et n'annonce plus de suite",
         p2.get("rangs") == "13 à 22 sur 22" and p2.get("pour_continuer") is None, str(p2.get("rangs")))

outils = (BACKEND / "skills" / "outils.py").read_text(encoding="utf-8")
verifier("le geste accepte `page` et la consigne interdit de remontrer la même page",
         '"page"]' in outils and "ne remontre" in outils and "page=data.get(\"page\") or 1" in outils)

print(f"\n{'═' * 70}\n{'✗ ' + str(len(echecs)) + ' échec(s) : ' + ', '.join(echecs) if echecs else '✓ 0 échec'}\n")
sys.exit(1 if echecs else 0)
