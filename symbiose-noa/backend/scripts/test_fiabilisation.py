"""Cas exécutables : concurrence interprocessus, versions, images, sources et délais."""
import asyncio,io,json,os,sys,tempfile,subprocess,types
from pathlib import Path
from unittest.mock import patch
B=Path(sys.argv[1]).resolve();sys.path.insert(0,str(B))
os.environ['DOCUMENTS_DIR']=tempfile.mkdtemp(prefix='fiabilisation-test-')
from bureautique import atelier,images
from PIL import Image
from docx import Document
from docx.shared import Inches

def image(couleur):
 b=io.BytesIO();Image.new('RGB',(24,24),couleur).save(b,format='PNG');return b.getvalue()

def concurrents(code,n=6):
 env=dict(os.environ,PYTHONPATH=str(B),PYTHONDONTWRITEBYTECODE='1')
 ps=[subprocess.Popen([sys.executable,'-c',code,str(i)],env=env,stdout=subprocess.PIPE,stderr=subprocess.PIPE,text=True) for i in range(n)]
 for p in ps:
  a,e=p.communicate(timeout=40);assert p.returncode==0,e

from ressources import registre
registre.noter('piece','ancienne',{'boite':'test@example.invalid'})
concurrents('''import sys
from ressources import registre
for j in range(25): registre.noter('piece', sys.argv[1]+':'+str(j), {'identifiant':str(j)})
''')
assert registre.combien('piece')==151 # le cache doit voir les autres processus
assert registre.lire('piece','ancienne')['boite']=='test@example.invalid'
print('✓ 150 écritures interprocessus sans perte, cache actualisé')
jeton=atelier.ouvrir({'titre':'Essai','format':'docx'},'personne')
concurrents('''import sys
from bureautique import atelier
for j in range(5): atelier.ajouter('''+repr(jeton)+''', [{'bloc':'texte','texte':'Paragraphe '+sys.argv[1]+'-'+str(j)}], 'personne',refuser_repetition=False)
''')
assert atelier.fiche(jeton,'personne')['elements']==30
img=atelier.ranger_image(jeton,'personne',image('red'),'png')
atelier.ajouter(jeton,[{'bloc':'image','fichier':img,'legende':'Image de test'}],'personne',False)
f=atelier.terminer(jeton,'personne');fichier=Path(atelier.chemin_fichier(jeton,'personne'));avant=fichier.read_bytes()
assert atelier.terminer(jeton,'personne')['empreinte']==f['empreinte'] and fichier.read_bytes()==avant
nouveau=atelier.nouvelle_revision(jeton,'personne');assert len(list(atelier.elements(nouveau)))==31
nouvelle_image=[e['fichier'] for e in atelier.elements(nouveau) if e.get('fichier')][0]
assert nouvelle_image.startswith(nouveau+'.img')
atelier.abandonner(jeton,'personne');assert Path(atelier.chemin_image(nouvelle_image)).exists()
assert atelier.terminer(nouveau,'personne')['fini']
print('✓ 30 versements concurrents, rendu immuable, révision avec contenu et images autonomes')
d=Document();d.sections[0].header.paragraphs[0].add_run().add_picture(io.BytesIO(image('blue')),width=Inches(.5));d.add_paragraph('Corps')
buf=io.BytesIO();d.save(buf)
png,mime,nom=images.image_du_docx(buf.getvalue(),'modele.docx');assert png==image('blue')
d.sections[0].header.paragraphs[0].add_run().add_picture(io.BytesIO(image('red')),width=Inches(.5));buf=io.BytesIO();d.save(buf)
try: images.image_du_docx(buf.getvalue(),'modele.docx')
except images.ImageRefusee:pass
else:raise AssertionError('Images ambiguës acceptées')
assert images.image_du_docx(buf.getvalue(),'modele.docx',numero=2)[0]==image('red')
print('✓ Logo Word incorporé, ambiguïté refusée, sélection explicite')
from skills.recherche_sources import completer
faux=types.ModuleType('skills.executor');appels=[]
async def execute(nom,args,**kw):
 appels.append((nom,kw['user'].id));return {'ok':True,'outcome':'success','output':{'resultats':[{'nom':'devis.docx'}]}}
faux.execute_skill=execute
async def recherches():
 user=types.SimpleNamespace(id='personne')
 with patch.dict(sys.modules,{'skills.executor':faux}):
  r=await completer({'nombre':0},{'requete':'devis'},user,'nas')
  assert len(r['recherche_sources'])==2 and set(appels)=={('nas_chercher','personne'),('nas_ouvrir','personne'),('lire_mails','personne')} and len(r['recherche_sources'][0]['lectures'])==1
  appels.clear();await completer({}, {'requete':'devis','sources_directes':False},user,'drive');assert not appels
  await completer({}, {'requete':'devis','types':['email']},user,'drive');assert appels==[('lire_mails','personne')]
asyncio.run(recherches());print('✓ Recherche de secours, identité transmise, filtre de source et désactivation')
from llm.budget import borner_tour
config=types.ModuleType('config');config.settings=types.SimpleNamespace(demande_delai_s=1)
@borner_tour
async def lent(): await asyncio.sleep(3)
with patch.dict(sys.modules,{'config':config}):
 try:asyncio.run(lent())
 except TimeoutError:pass
 else:raise AssertionError('Budget dépassé')
print('✓ Le budget interrompt réellement un tour trop long')

# Pièces : la vérification porte sur les octets, pas seulement sur le nom.
from mail.instantanes import verifier as verifier_pieces
import hashlib
piece={'nom':'devis.pdf','octets':b'contenu approuve'}
args={'_pieces_figees':[{'nom':piece['nom'],'sha256':hashlib.sha256(piece['octets']).hexdigest()}]}
verifier_pieces(args,[piece])
try: verifier_pieces(args,[{**piece,'octets':b'autre version'}])
except ValueError:pass
else:raise AssertionError('Pièce modifiée acceptée')
print('✓ Une pièce modifiée après accord bloque l’envoi')
# Pas de faux apprentissage après un effet ambigu ou une panne d’outil.
from learning.lecons import lire_lecon,_panne_passagere
lecon=lire_lecon(json.dumps({'lecon':{'situation':'Quand un devis doit être envoyé','erreur':'Erreur de fichier','conduite':'Vérifier la pièce avant de demander un accord','confiance':'invalide','type':'procedure'}}))
assert lecon['confiance']==.6
assert _panne_passagere({'tool_results':[{'outcome':'failed','error':'panne'}]})
assert _panne_passagere({'tool_results':[{'effect_status':'unknown'}]})
print('✓ Confiance invalide tolérée ; aucun apprentissage tiré d’un effet inconnu')

# Les octets approuvés survivent à la disparition de la source d'atelier.
from mail.instantanes import figer
from mail.attaches import resoudre
async def instantane_reel():
 user=types.SimpleNamespace(id='personne',role='admin')
 source=atelier.deposer_fichier('devis.pdf',b'%PDF-1.4 devis approuve','personne')
 fake_skills=types.ModuleType('mail.skills');fake_auth=types.ModuleType('mail.authorization')
 async def boite(data,user):return 'test@example.invalid'
 async def acces(user,mailbox,**kwargs):return mailbox
 fake_skills._boite_a_lire=boite;fake_auth.verifier_acces=acces
 with patch.dict(sys.modules,{'mail.skills':fake_skills,'mail.authorization':fake_auth}):
  accord=await figer({'pieces':[{'ref':source,'nom':'Devis-client.pdf'}]},user)
 assert accord['pieces'][0]['ref']!=source
 atelier.abandonner(source,'personne')
 pieces,refus=await resoudre(accord['pieces'],user,accord['mailbox'])
 assert not refus and pieces[0]['octets']==b'%PDF-1.4 devis approuve'
 verifier_pieces(accord,pieces)
 assert pieces[0]['nom']=='Devis-client.pdf'
asyncio.run(instantane_reel())
print('✓ Copie de pièce réelle : source supprimée, octets et nom approuvés conservés')

# Une panne disque ne doit pas annoncer une référence durable inexistante.
with patch.object(registre.os,'replace',side_effect=OSError('disque indisponible')):
 try:registre.noter('piece','ecriture-refusee',{'identifiant':'refusee'})
 except OSError:pass
 else:raise AssertionError('Écriture de référence perdue silencieusement')
assert registre.lire('piece','ecriture-refusee') is None
assert registre.lire('piece','ancienne') is not None
print('✓ Échec disque explicite ; registre précédent conservé')
registre.lire('piece','ancienne')
with patch.object(registre.json,'loads',side_effect=AssertionError('Relecture JSON inutile')):
 assert registre.lire('piece','ancienne') is not None
print('✓ Les lectures inchangées utilisent le cache vérifié par les métadonnées du fichier')

# Sans variable explicite, les trois entrepôts doivent désigner le même projet.
import ast
defauts=[]
for f in ['ressources/registre.py','bureautique/atelier.py','visuels/depot.py']:
 arbre=ast.parse((B/f).read_text())
 defauts.extend(n.args[1].value for n in ast.walk(arbre)
  if isinstance(n,ast.Call) and len(n.args)>1 and isinstance(n.args[0],ast.Constant)
  and n.args[0].value=='DOCUMENTS_DIR' and isinstance(n.args[1],ast.Constant))
assert len(defauts)==3 and len(set(defauts))==1,defauts
print('✓ Registre, atelier et visuels utilisent le même stockage par défaut')
print('✓ 0 échec')
