"""
Banc des DOCUMENTS QUI NE SE RESSEMBLENT PLUS — et du logo qui y entre enfin (15/09).

Relevé de Noa (Symbiose) : « il est toujours incapable d'inclure dans des
documents Word les images qu'il récupère, et il fait tout le temps la même
trame de documents. C'est horrible. Il doit absolument être capable de mettre
des en-têtes, des images qu'il récupère, etc. » Lu dans l'export du 15/09 :
  · 17:44 — `entete_image: "SYMBIOSE Paysage coul.pdf"` : un logo VECTORIEL,
    sans image intégrée → « ne porte aucune image exploitable », pas de logo ;
  · 17:56 et 18:43 — « récupère l'en-tête et le pied de ce devis » : un devis
    SCANNÉ → refusé (règle du 10/09 : « mieux vaut pas de logo qu'un scan ») ;
  · 18:43 — le logo PNG trouvé APRÈS la production : le second
    `produire_document` qui l'ajoutait était refusé (« un document fini ne se
    rallonge pas »), et la réponse affirmait « logo en en-tête de chaque page » ;
  · « ._SYMBIOSE-Paysage_rvb.png » (176 octets de métadonnées macOS) listé en
    tête, et « SYMBIOSE-Paysage_rvb.png » ouvrant « …-reserve.png » ;
  · chaque Word sortait avec la même page de garde et les mêmes titres.

CE QUE CE BANC PROUVE (les rendus sont RÉELS, puis rouverts) :
  · un logo PDF vectoriel est dessiné et recadré ; sur un devis scanné, la
    bande du HAUT (en-tête) ou du BAS (pied) est prise ;
  · quatre styles qui se voient (bandeau de couverture, couverture photo, sans
    page de garde…), et les blocs encadré, chiffres, citation, colonnes (photo à
    côté du texte) — en Word, en PDF et en Excel, sans casser le schéma Word ;
  · la présentation d'un document ouvert se change par `ajouter_document`, et
    refaire un document AVEC un logo n'est plus refusé ;
  · les fichiers parasites du Drive ne sont ni listés ni ouverts, le nom exact
    passe devant.
Exige python-docx, reportlab, openpyxl, PyMuPDF et Pillow pour les rendus
(sinon le banc le DIT et saute ces parties).

Usage : python backend/scripts/test_documents_varies.py [backend]
"""
import asyncio
import io
import os
import pathlib
import sys
import tempfile
import types
import zipfile

racine = pathlib.Path(sys.argv[1] if len(sys.argv) > 1 else "backend").resolve()
sys.path.insert(0, str(racine))
os.environ["DOCUMENTS_DIR"] = tempfile.mkdtemp(prefix="banc-docs-")
echecs = []


def verifier(nom, cond, detail=""):
    print(f"  {'✓' if cond else '✗'} {nom}" + (f"  → {str(detail)[:300]}" if detail and not cond else ""))
    if not cond:
        echecs.append(nom)


sys.modules.setdefault("emails", types.ModuleType("emails"))
sys.modules["emails.marque"] = types.SimpleNamespace(MARQUE={"couleur": "#2F6B3A", "fond": "#1E3D24", "nom": "Maison"})
from bureautique import modele  # noqa: E402

print("1. Le vocabulaire")
e = modele.normaliser_entete({"titre": "Offre", "style": "Moderne", "image_couverture": "chantier.jpg"})
verifier("le style est lu (vocabulaire fermé) et la couverture gardée",
         e["style"] == "moderne" and e["image_couverture"] == "chantier.jpg")
verifier("un style inventé retombe sur « classique »", modele.normaliser_entete({"style": "baroque"})["style"] == "classique")
n = modele.normaliser_element
verifier("encadré, citation, chiffres, colonnes sont des blocs",
         n({"bloc": "encadre", "texte": "Garantie décennale", "ton": "attention"})["ton"] == "attention"
         and n({"type": "quote", "text": "Très satisfaits", "auteur": "M. Camp"})["bloc"] == "citation"
         and n({"bloc": "chiffres", "items": [{"valeur": "450 m²", "libelle": "surface"}, ["6 sem.", "délai"]]})["items"][1]["valeur"] == "6 sem."
         and n({"bloc": "colonnes", "texte": "La terrasse", "image": "photo.jpg"})["ref"] == "photo.jpg")

try:
    import docx  # noqa: F401
    import fitz  # noqa: F401
    import openpyxl  # noqa: F401
    import reportlab  # noqa: F401
    from PIL import Image, ImageDraw
    rendus = True
except ImportError as err:
    print(f"\n  (bibliothèques de rendu absentes — {err} : parties 2 à 4 sautées ; le conteneur les a)")
    rendus = False

if rendus:
    from bureautique import atelier, images, rendu
    from docx import Document

    print("2. Un logo PDF, un devis scanné")
    # Un logo VECTORIEL : des tracés, aucune image intégrée.
    pdf = fitz.open()
    page = pdf.new_page(width=595, height=842)
    page.draw_rect(fitz.Rect(200, 380, 395, 460), color=(0.18, 0.42, 0.23), fill=(0.18, 0.42, 0.23))
    page.insert_text((215, 430), "SYMBIOSE", fontsize=26, color=(1, 1, 1))
    logo_pdf = pdf.tobytes()
    pdf.close()
    verifier("le PDF du logo ne porte aucune image intégrée (le cas de prod)",
             not fitz.open(stream=logo_pdf, filetype="pdf")[0].get_images())
    octets, mime, nom = images.logo_du_pdf(logo_pdf, "SYMBIOSE Paysage coul.pdf")
    img = Image.open(io.BytesIO(octets))
    verifier("il est DESSINÉ et recadré sur le logo (plus « aucune image exploitable »)",
             mime == "image/png" and img.width < 700 and img.height < 300 and "dessiné" in nom, (img.size, nom))

    # Un devis SCANNÉ : une seule image pleine page, logo en haut, mentions en bas.
    scan = Image.new("RGB", (1240, 1754), "white")
    d = ImageDraw.Draw(scan)
    d.rectangle((80, 60, 520, 220), fill=(40, 100, 50))                  # en-tête : logo
    for y in range(400, 1400, 40):
        d.rectangle((80, y, 1160, y + 12), fill=(90, 90, 90))            # corps du devis
    d.rectangle((80, 1640, 1160, 1690), fill=(120, 120, 120))            # pied : mentions
    tampon = io.BytesIO()
    scan.save(tampon, format="PNG")
    pdf = fitz.open()
    page = pdf.new_page(width=595, height=842)
    page.insert_image(page.rect, stream=tampon.getvalue())
    devis_pdf = pdf.tobytes()
    pdf.close()
    haut, _, nom_h = images.logo_du_pdf(devis_pdf, "DV0001224.pdf")
    bas, _, nom_b = images.logo_du_pdf(devis_pdf, "DV0001224.pdf", "pied")
    ih, ib = Image.open(io.BytesIO(haut)), Image.open(io.BytesIO(bas))
    verifier("devis scanné : l'EN-TÊTE est la bande du haut (le logo), pas la page entière",
             "en-tête" in nom_h and ih.height < ih.width and ih.height < 450, (ih.size, nom_h))
    verifier("… et le PIED est la bande du bas", "pied" in nom_b and ib.height < 200 and ib.width > 1000, (ib.size, nom_b))

    print("3. Des documents qui ne se ressemblent pas")
    proprio = "u-banc"
    photo = Image.new("RGB", (800, 500), (120, 160, 90))
    tp = io.BytesIO(); photo.save(tp, format="JPEG")

    def produire(style, blocs, avec_couverture=True):
        entete = modele.normaliser_entete({"titre": "Aménagement Camp", "sous_titre": "Proposition",
                                           "style": style, "format": "docx", "pied": "Maison · 05 56"})
        jeton = atelier.ouvrir(entete, proprio)
        entete["entete_image_fichier"] = atelier.ranger_image(jeton, proprio, octets, "png")
        if avec_couverture:
            entete["image_couverture_fichier"] = atelier.ranger_image(jeton, proprio, tp.getvalue(), "jpg")
        fichier_photo = atelier.ranger_image(jeton, proprio, tp.getvalue(), "jpg")
        elements = [x for x in (modele.normaliser_element(b) for b in blocs) if x]
        for x in elements:
            if x["bloc"] in ("colonnes", "image") and x.get("ref"):
                x["fichier"] = fichier_photo
        chemin = os.path.join(tempfile.mkdtemp(), f"{style}.docx")
        rendu._docx(entete, elements, chemin)
        return chemin

    blocs = [{"bloc": "titre", "texte": "Le projet", "niveau": 1},
             {"bloc": "paragraphe", "texte": "Une terrasse en pin traité, des massifs de graminées."},
             {"bloc": "chiffres", "items": [{"valeur": "120 m²", "libelle": "terrasse"},
                                            {"valeur": "6 semaines", "libelle": "chantier"}]},
             {"bloc": "colonnes", "titre": "La terrasse", "texte": "Pin classe 4.\n- lames 27 mm", "image": "photo"},
             {"bloc": "encadre", "titre": "Garantie", "texte": "Décennale sur la maçonnerie.", "ton": "attention"},
             {"bloc": "citation", "texte": "Un travail soigné", "auteur": "M. Camp"}]
    xmls = {}
    for style in ("classique", "moderne", "epure", "plaquette"):
        chemin = produire(style, blocs)
        doc = Document(chemin)                              # se rouvre : le fichier est sain
        with zipfile.ZipFile(chemin) as z:
            xmls[style] = z.read("word/document.xml").decode("utf-8")
        verifier(f"« {style} » : le Word se rouvre et porte la photo À CÔTÉ du texte (colonnes)",
                 len(doc.tables) >= 3 and "La terrasse" in xmls[style] and xmls[style].count("<pic:pic") >= 2)
    verifier("« moderne » : un bandeau de couverture à la couleur de la maison",
             'w:fill="2F6B3A"' in xmls["moderne"].split("Le projet")[0])
    verifier("« epure » : pas de page de garde (aucun saut de page avant le contenu)",
             'w:type="page"' not in xmls["epure"].split("Le projet")[0])
    verifier("« plaquette » : la photo de couverture avant le titre",
             xmls["plaquette"].index("<pic:pic") < xmls["plaquette"].index("Aménagement Camp"))
    verifier("les quatre Word ne sont pas le même document",
             len({x.split("Le projet")[0][:4000] for x in xmls.values()}) == 4)
    verifier("chiffres, encadré et citation sont écrits",
             "120 m²" in xmls["moderne"] and "Décennale" in xmls["moderne"] and "Un travail soigné" in xmls["moderne"])
    import re
    ordre_ok = all(
        not re.search(r"<w:(?:spacing|ind|jc)\b[^>]*/>(?:(?!</w:pPr>).)*<w:pBdr>", x)
        and not re.search(r"<w:shd\b[^>]*/>(?:(?!</w:tcPr>).)*<w:tcBorders>", x)
        for x in xmls.values())
    verifier("filets et bordures À LEUR PLACE dans le schéma (sinon Word : « contenu illisible »)", ordre_ok)

    print("4. Le PDF et l'Excel suivent")
    entete_pdf = modele.normaliser_entete({"titre": "Offre", "format": "pdf", "style": "moderne"})
    elements = [x for x in (modele.normaliser_element(b) for b in blocs) if x]
    for x in elements:
        x.pop("ref", None)
    sortie_pdf = os.path.join(tempfile.mkdtemp(), "offre.pdf")
    rendu._pdf(entete_pdf, elements, sortie_pdf)
    texte_pdf = "".join(p.get_text() for p in fitz.open(sortie_pdf))
    verifier("PDF : chiffres, encadré, citation rendus", "120 m²" in texte_pdf and "Décennale" in texte_pdf
             and "Un travail soigné" in texte_pdf, texte_pdf[:200])
    sortie_x = os.path.join(tempfile.mkdtemp(), "offre.xlsx")
    rendu._xlsx(modele.normaliser_entete({"titre": "Offre", "format": "xlsx"}), elements, sortie_x)
    valeurs = [c.value for ws in openpyxl.load_workbook(sortie_x).worksheets for r in ws.iter_rows() for c in r]
    verifier("Excel : les chiffres et l'encadré y sont, en cellules", "120 m²" in valeurs
             and any("Décennale" in str(v) for v in valeurs))

print("5. Le logo trouvé après coup, et les fichiers parasites")
bureau = (racine / "skills" / "bureau.py").read_text(encoding="utf-8")
verifier("`ajouter_document` change la présentation d'un document ouvert (logo, couverture, style)",
         '"entete_image", "pied_image", "image_couverture", "style"' in bureau and "presentation_modifiee" in bureau)
docs = (racine / "outils" / "documents.py").read_text(encoding="utf-8")
verifier("refaire le même titre AVEC un logo n'est plus refusé comme un « rallongement »",
         "nouvelle_presentation" in docs and "and not nouvelle_presentation" in docs)
proto = (racine / "skills" / "protocol.py").read_text(encoding="utf-8")
verifier("le catalogue fait CHOISIR le style et varier la page", "CHOISIS-LE" in proto and "VARIE la page" in proto)
drive = racine / "outils" / "drive.py"
if drive.exists():
    src = drive.read_text(encoding="utf-8")
    esp = {}
    import ast
    arbre = ast.parse(src)
    exec(compile(ast.Module(body=[x for x in arbre.body if isinstance(x, ast.FunctionDef) and x.name in ("_parasite", "_classer")]
                            + [x for x in arbre.body if isinstance(x, ast.Assign) and any(getattr(t, "id", "") == "_MIME_DOSSIER" for t in x.targets)],
                            type_ignores=[]), "drive", "exec"), esp)
    d, f = esp["_classer"]([{"name": "._SYMBIOSE-Paysage_rvb.png"}, {"name": "SYMBIOSE-Paysage_rvb.png"},
                            {"name": "__MACOSX", "mimeType": esp["_MIME_DOSSIER"]}, {"name": ".DS_Store"}])
    verifier("Drive : « ._logo.png », « __MACOSX », « .DS_Store » ne sont ni listés ni comptés",
             not d and [x["name"] for x in f] == ["SYMBIOSE-Paysage_rvb.png"], (d, f))
    verifier("Drive : à l'ouverture, le nom EXACT passe devant « …-reserve.png »",
             "0 if _nu(f.get(\"name\") or \"\") == _nu(nom) else 1" in src)
else:
    print("  (pas de Drive ici : contrôles du Drive sautés)")

print()
if echecs:
    print(f"✗ {len(echecs)} échec(s)")
    sys.exit(1)
print("✓ 0 échec")
