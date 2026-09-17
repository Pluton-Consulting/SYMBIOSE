"""Recette locale des dossiers longs : fichiers réels, stockage isolé, modèles simulés.

Les modèles simulés vérifient l'orchestration et les reprises, pas la qualité
sémantique d'un fournisseur. Le rejeu du texte Langfuse est optionnel via EXPORT_LANGFUSE.
"""
import ast,asyncio,io,json,os,pathlib,sys,tempfile,unittest
from types import SimpleNamespace
from unittest.mock import patch,AsyncMock
BACKEND=pathlib.Path(sys.argv[1] if len(sys.argv)>1 else 'backend').resolve();sys.path.insert(0,str(BACKEND));sys.argv=[sys.argv[0]]
TEMP=tempfile.TemporaryDirectory(prefix='banc-dossiers-');os.environ['DOCUMENTS_DIR']=TEMP.name+'/documents'
from ressources import dossiers,registre
from bureautique import atelier
from bureautique.lecture_integrale import lire
from security.conversation import fil_courant
from skills import bureau,documents_dossier

def fonctions(path,noms,env=None):
    tree=ast.parse(path.read_text());nodes=[n for n in tree.body if isinstance(n,(ast.FunctionDef,ast.AsyncFunctionDef)) and n.name in noms]
    ns={'AgentState':dict,**(env or {})};exec(compile(ast.Module(body=nodes,type_ignores=[]),str(path),'exec'),ns);return ns

class Recette(unittest.TestCase):
 def test_couverture_ne_reprend_pas_un_remplacement_de_pied(self):
    contrat={'titre':'Mémoire technique — Chantier actuel'}
    plan={'titre':'p. 1 Entreprise — Confidentiel','remplacements_modele':{'p. 1 Confidentiel':'p.  1 Entreprise — Confidentiel'}}
    self.assertEqual(documents_dossier._titre_livrable(plan,contrat),contrat['titre'])
    self.assertEqual(documents_dossier._titre_livrable({**plan,'titre':'Rapport technique'},contrat),'Rapport technique')
    self.assertEqual(documents_dossier._titre_livrable({'titre':'Rapport technique'},contrat),'Rapport technique')
 def test_controle_factuel_partitionne_sans_perte_et_repris(self):
    analyses=[{'source':'p','preuve':'p:1','faits':[{'fait':'Fait '+str(i)+' x'*80,'nature':'exigence','citation':'source'} for i in range(25)],'limites':[]}]
    paquets=documents_dossier._paquets_preuves(analyses,800)
    self.assertGreater(len(paquets),1)
    self.assertEqual([f['fait'] for p in paquets for a in p for f in a['faits']],[f['fait'] for f in analyses[0]['faits']])
    appels=[]
    async def modele(consigne,data,verifier=None,**options):
     appels.append(data)
     r={'problemes':[{'section':0,'raison':'Date contredite','sources':['p:1']}]}
     if verifier:verifier(r)
     return r
    async def test():
     args=(self.uid,self.fil,'tache',{'sections':[{'titre':'Planning'}]},[{'blocs':[{'bloc':'paragraphe','texte':'Date'}],'reserves':[]}],analyses)
     with patch.object(documents_dossier,'_json',modele):
      premier=await documents_dossier._controler_faits(*args)
      n=len(appels);second=await documents_dossier._controler_faits(*args)
     self.assertEqual(premier,second);self.assertEqual(len(appels),n)
     self.assertEqual(premier[0]['sources'],['p:1'])
    asyncio.run(test())
 def test_arbitrage_rejette_fausse_contradiction_et_recontrole_modification(self):
    analyses=[{'source':'p','preuve':'p:1','faits':[{'fait':'Pénalité A : 100 euros ; pénalité B : 150 euros.'}]}]
    plan={'sections':[{'titre':'Pénalités','sources':['p']}]}
    sections=[{'blocs':[{'bloc':'paragraphe','texte':'Pénalité B : 150 euros.'}],'preuves':['p:1'],'reserves':[]}]
    appels=[]
    async def modele(consigne,data,verifier=None,**options):
     appels.append(consigne)
     if consigne.startswith('Arbitre'):
      self.assertIn('Pénalité A',data['preuves_completes_concernees'][0]['faits'][0]['fait'])
      r={'problemes':[]}
     else:r={'problemes':[{'section':0,'raison':'Toute pénalité serait de 100 euros','sources':['p:1']}]}
     if verifier:verifier(r)
     return r
    async def test():
     with patch.object(documents_dossier,'_json',modele):
      self.assertEqual(await documents_dossier._controler_faits(self.uid,self.fil,'a',plan,sections,analyses),[])
      self.assertEqual(len(appels),2)
      self.assertEqual(await documents_dossier._controler_faits(self.uid,self.fil,'a',plan,sections,analyses),[])
      self.assertEqual(len(appels),2)
      sections[0]['blocs'][0]['texte']='Pénalité B : 100 euros.'
      await documents_dossier._controler_faits(self.uid,self.fil,'a',plan,sections,analyses)
      self.assertEqual(len(appels),4)
    asyncio.run(test())
 def test_arbitrage_refuse_preuve_inventee(self):
    async def modele(consigne,data,verifier=None,**options):
     r={'problemes':[{'section':0,'raison':'Contradiction','sources':['etrangere' if consigne.startswith('Arbitre') else 'p:1']}]}
     if verifier:verifier(r)
     return r
    async def test():
     with patch.object(documents_dossier,'_json',modele):
      with self.assertRaisesRegex(ValueError,'Preuves arbitrées'):
       await documents_dossier._controler_faits(self.uid,self.fil,'a',{'sections':[{'titre':'S','sources':['p']}]},[{'blocs':[],'reserves':[]}],[{'source':'p','preuve':'p:1','faits':[{'fait':'F'}]}])
    asyncio.run(test())
 def test_correction_factuelle_ne_melange_pas_les_sections(self):
    c=documents_dossier._correction_factuelle([
     {'section':0,'raison':'Adresse','sources':['a:1']},
     {'section':2,'raison':'Planning','sources':['p:2']},
     {'section':2,'raison':'Durée','sources':['p:2','c:1']}])
    self.assertEqual(c['cibles'],[0,2]);self.assertTrue(c['arbitre'])
    self.assertEqual(c['par_section']['0'],{'problemes':['Adresse'],'preuves_complementaires':['a:1']})
    self.assertEqual(c['par_section']['2'],{'problemes':['Planning','Durée'],'preuves_complementaires':['c:1','p:2']})
 def test_reserves_par_lots_reprend_seulement_le_lot_echoue(self):
    analyses=[{'source':'p','preuve':'p:1','faits':[{'fait':str(i)+' information'*35} for i in range(300)]}]
    appels=[];echouer=True
    async def modele(consigne,data,verifier=None,**options):
     nonlocal echouer
     ident=data['faits_dossier'][0]['faits'][0];appels.append(ident)
     self.assertLess(len(json.dumps(data,ensure_ascii=False)),46000)
     if echouer and ident.startswith('0 '):
      echouer=False;raise TimeoutError('simulé')
     resultat={'problemes':[]}
     if verifier:verifier(resultat)
     return resultat
    async def test():
     args=({},[{'id':'p','nom':'Pièce'}],[{'reserves':['Effectif à confirmer']}],analyses,(self.uid,self.fil,'reserves'))
     with patch.object(documents_dossier,'_json',modele):
      with self.assertRaisesRegex(ValueError,'partiel conservé'):
       await documents_dossier._controler_reserves(*args)
      n=len(appels);self.assertGreater(n,1)
      self.assertEqual(await documents_dossier._controler_reserves(*args),{'problemes':[]})
      self.assertEqual(len(appels),n+1)
      await documents_dossier._controler_reserves(*args)
      self.assertEqual(len(appels),n+1)
    asyncio.run(test())
 def test_reserve_arbitree_avec_la_piece_citee_dans_son_texte(self):
    a='a'*24;b='b'*24
    reserve='Le cadre et la notice divergent ('+b+':1).'
    analyses=[{'source':a,'preuve':a+':1','faits':[{'fait':'Adresse du cadre A'}]},
              {'source':b,'preuve':b+':1','faits':[{'fait':'Adresse de la notice B'}]}]
    appels=[]
    async def modele(consigne,data,verifier=None,**options):
     appels.append(consigne)
     if consigne.startswith('Évalue directement'):
      self.assertEqual({a+':1',b+':1'},{p['preuve'] for p in data['preuves_completes_concernees']})
      r={'avis':[{'reserve':reserve,'fondee':True,'raison':'Les deux pièces divergent réellement.','sources':[a+':1',b+':1']}]}
     else:r={'problemes':[{'section':0,'reserve':reserve,'raison':'Le cadre donne A','sources':[a+':1']}]}
     if verifier:verifier(r)
     return r
    async def test():
     with patch.object(documents_dossier,'_json',modele):
      r=await documents_dossier._controler_reserves({'sections':[{'titre':'Coordonnées','sources':[a]}]},
       [{'id':a,'nom':'Cadre'},{'id':b,'nom':'Notice'}],
       [{'blocs':[{'bloc':'paragraphe','texte':'Coordonnées conservées.'}],'preuves':[a+':1'],'reserves':[reserve]}],analyses,(self.uid,self.fil,'conflit'))
     self.assertEqual(r,{'problemes':[]});self.assertEqual(len(appels),2)
    asyncio.run(test())
 def test_arbitrage_reserve_retire_faux_conflit_et_exige_un_verdict_par_reserve(self):
    reserve='Conflit : annexe 1 du RC différente de celle du CCAP.'
    analyses=[{'source':'a','preuve':'a:1','faits':[{'fait':'Annexe 1 du RC : attestation.'}]},
              {'source':'b','preuve':'b:1','faits':[{'fait':'Annexe 1 du CCAP : insertion.'}]}]
    async def test():
     async def modele(consigne,data,verifier=None,**options):
      self.assertEqual(data['reserves_a_evaluer'],[reserve])
      r={'avis':[{'reserve':reserve,'fondee':False,'raison':'Deux annexes de pièces différentes, sans conflit démontré.','sources':['a:1','b:1']}]}
      if verifier:verifier(r)
      return r
     args=(self.uid,self.fil,'reserve-fausse',{'sections':[{'titre':'Candidature','sources':['a','b']}]},[{'blocs':[],'reserves':[reserve]}],analyses,[{'section':0,'reserve':reserve,'raison':'Fausse réserve','sources':['a:1','b:1']}])
     with patch.object(documents_dossier,'_json',modele):
      r=await documents_dossier._arbitrer_faits(*args,reserves=True)
     self.assertEqual(r[0]['reserve'],reserve);self.assertEqual(r[0]['section'],0)
     async def incomplet(consigne,data,verifier=None,**options):
      r={'avis':[]}
      if verifier:verifier(r)
      return r
     with patch.object(documents_dossier,'_json',incomplet):
      with self.assertRaisesRegex(ValueError,'Chaque réserve'):await documents_dossier._arbitrer_faits(*(*args[:2],'autre',*args[3:]),reserves=True)
    asyncio.run(test())
 def test_reprise_accepte_corrections_structurees_sans_trier_des_objets(self):
    anciens=[{'raison':'Ancienne alerte'}]
    self.assertFalse(documents_dossier._memes_signalements(anciens,[{'rubrique':'Planning','description':'Durée'}, {'rubrique':'Titre','description':'Référence'}]))
    self.assertTrue(documents_dossier._memes_signalements(anciens,['Ancienne alerte']))
    self.assertFalse(documents_dossier._memes_signalements(anciens,['Nouvelle alerte']))
 def test_paquets_mesurent_metadonnees_une_fois_et_bornent_la_taille(self):
    analyses=[{'source':'p','preuve':'p:1','limites':['Contexte '*70],
               'faits':[{'fait':'Information '+str(i)} for i in range(100)]}]
    paquets=documents_dossier._paquets_preuves(analyses,2000)
    self.assertLess(len(paquets),5)
    self.assertTrue(all(len(json.dumps(p,ensure_ascii=False))<=2000 for p in paquets))
    self.assertEqual([f for p in paquets for a in p for f in a['faits']],analyses[0]['faits'])
    self.assertEqual(len(analyses[0]['faits']),100)
    self.assertTrue(all(a['limites']==analyses[0]['limites'] for p in paquets for a in p))
 def test_synthese_garde_tous_les_faits_sans_repeter_les_citations(self):
    faits=[{'fait':'Condition '+str(i),'nature':'exigence','citation':'Tableau complet '*200,'lignes':[1,4]} for i in range(120)]
    source=[{'source':'p1','preuve':'p1:1','faits':faits,'limites':['Borne incertaine']}]
    sortie=documents_dossier._faits_pour_synthese(source)
    self.assertEqual(len(sortie[0]['faits']),120)
    self.assertEqual([f['fait'] for f in sortie[0]['faits']],[f['fait'] for f in faits])
    self.assertEqual(sortie[0]['limites'],['Borne incertaine'])
    self.assertEqual(sortie[0]['preuve'],'p1:1')
    self.assertTrue(all(f['nature']=='exigence' for f in sortie[0]['faits']))
    self.assertIn('citation',source[0]['faits'][0])
    self.assertNotIn('citation',sortie[0]['faits'][0])
 def test_reprise_par_identifiant_ne_cree_pas_deux_files(self):
    from ressources import documents_file as f
    source=dossiers.enregistrer(self.uid,self.fil,'a.txt','Pièce de référence')
    a=f.soumettre(self.uid,self.fil,'document',{'demande':'Rapport','sources':[source]})
    token=f._COURANTE.set(a['tache_documentaire'])
    try:f.associer_tache(self.uid,self.fil,'moteur-1')
    finally:f._COURANTE.reset(token)
    b=f.soumettre(self.uid,self.fil,'document',{'demande':'Reprends le rapport','tache':'moteur-1'})
    self.assertEqual(a['tache_documentaire'],b['tache_documentaire'])
    self.assertEqual(len(f.etats(self.uid,self.fil)),1)
    c=f.soumettre(self.uid,'autre-fil','document',{'demande':'Rapport','tache':'moteur-1'})
    self.assertNotEqual(c['tache_documentaire'],a['tache_documentaire'])
 def test_recherche_nas_noms_et_confirmation_superflue(self):
    if (BACKEND/'nas/acces.py').exists():
     ns=fonctions(BACKEND/'nas/acces.py',{'_nom_correspond','_sans_accent_nas'})
     correspond=ns['_nom_correspond']
     self.assertTrue(correspond('2029 - AIRBORNE_SONOVISION.pdf','2029 airborne sonovision'))
     self.assertTrue(correspond('17 IKOS CCTP 17 Plâtrerie.pdf','CCTP17'))
     self.assertFalse(correspond('CCTP 117.pdf','CCTP17'))
     self.assertFalse(correspond('Autre dossier.pdf','AIRBORNE'))
     self.assertFalse(correspond('document.pdf',''))
    from agents.annonce import propose_au_lieu_d_agir
    self.assertTrue(propose_au_lieu_d_agir("Je peux lister le dossier CCTP pour retrouver le fichier exact et l’ouvrir. Je le fais ?"))
    self.assertFalse(propose_au_lieu_d_agir("Quel dossier souhaitez-vous ouvrir ?"))
    self.assertTrue(propose_au_lieu_d_agir("La ligne est tronquée. Veux-tu que je la relise pour compléter, ou que je prépare des relances ?"))
    self.assertFalse(propose_au_lieu_d_agir("Où souhaitez-vous enregistrer le fichier ?"))
 def test_lecture_visuelle_suit_sa_piece_dans_toute_rubrique(self):
    plan={'sections':[{'sources':['planning','ccap']},{'sources':['ccap']},{'sources':['planning','page1']}], 'sources_ecartees':{'page1':'oubli','tiers':'sans rapport'}}
    sources=[{'id':'planning'},{'id':'page1','_source_originale':'planning'},{'id':'page2','_source_originale':'planning'},{'id':'tiers','_source_originale':'autre'}]
    self.assertEqual(documents_dossier._inclure_sources_derivees(plan,sources),[0,2])
    self.assertEqual(plan['sections'][0]['sources'],['planning','ccap','page1','page2'])
    self.assertEqual(plan['sections'][1]['sources'],['ccap'])
    self.assertEqual(plan['sources_ecartees'],{'tiers':'sans rapport'})
    self.assertEqual(documents_dossier._inclure_sources_derivees(plan,sources),[])
 def test_suivi_et_reprise_isoles_et_idempotents(self):
    from ressources import documents_file as f
    source=dossiers.enregistrer(self.uid,self.fil,'cadre.txt','Informations du cadre')
    job=f.soumettre(self.uid,self.fil,'document',{'demande':'Un rapport','sources':[source]})['tache_documentaire']
    f._maj(job,statut='en_cours',essais=4,prochain=123456)
    jeton=f._COURANTE.set(job)
    try:f.associer_tache(self.uid,self.fil,'tache-a')
    finally:f._COURANTE.reset(jeton)
    dossiers.etape(self.uid,self.fil,'tache-a','plan',{'sections':[{'titre':'A'},{'titre':'B'}]})
    dossiers.etape(self.uid,self.fil,'tache-a','section:0',{'texte':'Confidentiel'})
    dossiers.etape(self.uid,self.fil,'autre','section:1',{'texte':'Ne pas compter'})
    for identifiant in (job,'tache-a'):
     r=f.piloter(self.uid,self.fil,identifiant,True)
     self.assertTrue(r['en_cours']);self.assertEqual(r['tache_documentaire'],job)
    with dossiers.base() as c:
     r=c.execute('SELECT essais,prochain,statut FROM file_documentaire WHERE id=?',(job,)).fetchone()
    self.assertEqual(tuple(r),(4,123456,'en_cours'))
    suivi=f.progression(self.uid,self.fil)
    self.assertEqual(suivi[0]['sections'],1);self.assertEqual(suivi[0]['sections_total'],2)
    self.assertNotIn('Confidentiel',json.dumps(suivi))
    self.assertEqual(f.progression('autre',self.fil),[])
    self.assertEqual(f.progression(self.uid,'autre'),[])
    with self.assertRaises(ValueError):f.piloter('autre',self.fil,job,True)
    with self.assertRaises(ValueError):f.piloter(self.uid,'autre','tache-a',True)
    f.piloter(self.uid,self.fil,job,False)
    self.assertEqual(f.progression(self.uid,self.fil)[0]['statut'],'suspendu')
    self.assertTrue(f.piloter(self.uid,self.fil,job,True)['en_cours'])
    f._maj(job,statut='termine',annonce=0)
    self.assertFalse(f.progression(self.uid,self.fil)[0]['annonce'])
    f._maj(job,annonce=1)
    self.assertEqual(f.progression(self.uid,self.fil)[0]['phase'],'Document prêt')
 def test_references_affichees_resolues_sans_croiser_les_comptes(self):
    a=dossiers.enregistrer(self.uid,self.fil,'cadre.docx','Rubriques à remplir','/api/documents/cadre-prive')
    b=dossiers.enregistrer(self.uid,self.fil,'source.pdf','Exigences','/api/documents/source-prive')
    r=dossiers.normaliser_selection(self.uid,self.fil,{'sources':['/api/documents/cadre-prive',b], 'modele_source':'cadre.docx'})
    self.assertEqual(r,{'sources':[a,b],'modele_source':a})
    with self.assertRaises(ValueError):dossiers.sources('autre',self.fil,['/api/documents/cadre-prive'])
    with self.assertRaises(ValueError):dossiers.sources(self.uid,'autre',['/api/documents/cadre-prive'])
    dossiers.enregistrer(self.uid,self.fil,'cadre.docx','Autre cadre','/api/documents/autre-cadre')
    with self.assertRaises(ValueError):dossiers.sources(self.uid,self.fil,['cadre.docx'])
    self.assertEqual(dossiers.sources(self.uid,self.fil,[a])[0]['id'],a)
 def test_court_dossier_libere_chat_et_persiste(self):
    source=dossiers.enregistrer(self.uid,self.fil,'cadre.txt','Onze rubriques à remplir')
    async def test():
     with patch.object(documents_dossier,'composer_immediat',AsyncMock(side_effect=AssertionError('rédaction bloquante'))):
      a=await documents_dossier.composer({'_fil':self.fil,'demande':'Remplis ce cadre','sources':[source]},SimpleNamespace(id=self.uid))
      b=await documents_dossier.composer({'_fil':self.fil,'demande':'Remplis ce cadre','sources':[source]},SimpleNamespace(id=self.uid))
     self.assertTrue(a['en_cours']);self.assertEqual(a['tache_documentaire'],b['tache_documentaire'])
     from ressources.documents_file import etats
     self.assertEqual(len(etats(self.uid,self.fil)),1)
    asyncio.run(test())
 def test_formule_invalide_redevient_corrigeable(self):
    from skills.chiffres_sources import calculer
    self.assertEqual(calculer('c1*c2',{'c1':'5','c2':'4'}),'20')
    for expression in ('longueur x largeur','c1 *** c2','__import__("os")'):
     with self.assertRaises(ValueError):calculer(expression,{'c1':'5','c2':'4'})
 def setUp(self):
    self.rep=tempfile.TemporaryDirectory(dir=TEMP.name);self.addCleanup(self.rep.cleanup)
    self.patch=patch.object(registre,'_chemin',lambda:pathlib.Path(self.rep.name)/'registre.json');self.patch.start();self.addCleanup(self.patch.stop)
    self.uid='compte-test';self.fil='fil-1'
 def test_sept_pieces_integrales(self):
    contenus=['Document '+str(i)+'\n'+('Exigence importante.\n'*4500)+'FIN_SOURCE_'+str(i) for i in range(7)]
    ids=[dossiers.enregistrer(self.uid,self.fil,f'piece{i}.txt',t) for i,t in enumerate(contenus)]
    self.assertEqual(len(dossiers.manifeste(self.uid,self.fil)),7)
    for i,source in enumerate(ids):
     blocs=dossiers.fragments(contenus[i]);self.assertEqual(''.join(dossiers.lire(self.uid,self.fil,source,n+1)['texte'] for n in range(len(blocs))),contenus[i])
    self.assertEqual(dossiers.manifeste('autre',self.fil),[])
    with self.assertRaises(ValueError):dossiers.lire(self.uid,'autre-fil',ids[0])
 def test_homonymes_et_versions_reelles(self):
    a=dossiers.enregistrer(self.uid,self.fil,'CCTP.pdf','Version A','nas:/A/CCTP.pdf')
    b=dossiers.enregistrer(self.uid,self.fil,'CCTP.pdf','Version B','nas:/B/CCTP.pdf')
    self.assertEqual({r['id'] for r in dossiers.manifeste(self.uid,self.fil)},{a,b})
    c=dossiers.enregistrer(self.uid,self.fil,'CCTP.pdf','Version A corrigée','nas:/A/CCTP.pdf')
    self.assertEqual({r['id'] for r in dossiers.manifeste(self.uid,self.fil)},{c,b})
    self.assertEqual(dossiers.lire(self.uid,self.fil,a)['texte'],'Version A')
 def test_homonymes_sans_reference_et_contenu_identique(self):
    a=dossiers.enregistrer(self.uid,self.fil,'CCTP.pdf','A')
    b=dossiers.enregistrer(self.uid,self.fil,'CCTP.pdf','B')
    c=dossiers.enregistrer(self.uid,self.fil,'CCTP.pdf','A','nas:/C/CCTP.pdf')
    d=dossiers.enregistrer(self.uid,self.fil,'CCTP.pdf','A','nas:/D/CCTP.pdf')
    self.assertEqual(len({a,b,c,d}),4)
    self.assertEqual(len(dossiers.manifeste(self.uid,self.fil)),4)
    self.assertEqual(dossiers.enregistrer(self.uid,self.fil,'CCTP.pdf','A','nas:/C/CCTP.pdf'),c)
 def test_identifiants_historiques_conserves(self):
    import hashlib,time
    texte='Ancienne source';nom='CCTP.pdf';ref='nas:/ancien/CCTP.pdf'
    sha=hashlib.sha256(texte.encode()).hexdigest()
    cle=hashlib.sha256((self.uid+'\0'+self.fil+'\0'+nom+'\0'+sha+'\0'+texte).encode()).hexdigest()[:24]
    with dossiers.base() as c:c.execute('INSERT INTO sources_dossier VALUES(?,?,?,?,?,?,?,?)',(cle,self.uid,self.fil,nom,sha,ref,texte,time.time()))
    self.assertEqual(dossiers.enregistrer(self.uid,self.fil,nom,texte,ref),cle)
 def test_progression_ne_compte_que_la_redaction_concernee(self):
    # Exécuter l'expression SQL réellement employée par le worker.
    tree=ast.parse((BACKEND/'ressources/documents_file.py').read_text())
    sql=next(n.value for n in ast.walk(tree) if isinstance(n,ast.Constant) and isinstance(n.value,str) and n.value.startswith('SELECT count(*) FROM etapes_documentaires'))
    dossiers.etape(self.uid,self.fil,'ma-tache','contrat',{'demande':'rapport'})
    dossiers.etape(self.uid,self.fil,'autre-tache','section:1',{'ok':True})
    comptes=[]
    for i in range(3):
     dossiers.etape(self.uid,self.fil,'file:job','consommation:'+str(i),{'euros':0})
     with dossiers.base() as c:comptes.append(c.execute(sql,(self.uid,self.fil,'ma-tache')).fetchone()[0])
    self.assertEqual(comptes,[0,0,0])
    dossiers.etape(self.uid,self.fil,'ma-tache','analyse:source:1',{'faits':[]})
    with dossiers.base() as c:self.assertEqual(c.execute(sql,(self.uid,self.fil,'ma-tache')).fetchone()[0],1)
 def test_file_bloque_apres_echecs_sans_progression(self):
    from ressources import documents_file
    from contextvars import ContextVar
    from contextlib import nullcontext
    import types
    data={'demande':'Rapport sans source accessible'}
    travail=documents_file.soumettre(self.uid,self.fil,'document',data)
    cle=travail['tache_documentaire']
    faux={
     'tasks.identity':types.SimpleNamespace(charger_executant=AsyncMock(return_value=SimpleNamespace(id=self.uid,role='admin'))),
     'security.lecteur':types.SimpleNamespace(au_nom_de=lambda user:nullcontext()),
     'llm.budget':types.SimpleNamespace(Budget=lambda n:object(),_courant=ContextVar('budget-test',default=None)),
     'llm.concurrence':types.SimpleNamespace(PERSONNE=ContextVar('personne-test',default=None)),
     'llm.compteur':types.SimpleNamespace(_TOUR=ContextVar('tour-test',default=None),bilan=lambda:{'tokens_in':0,'tokens_out':0,'cost_eur':0}),
     'config':types.SimpleNamespace(settings=SimpleNamespace(llm_simultanes_fond=2))}
    async def erreur(*args):return {'tache':'tache-test','production_verifiee':False,'note':'Même échec sans progression'}
    async def test():
     with patch.dict(sys.modules,faux),patch.object(documents_dossier,'composer_immediat',erreur),patch.object(documents_file,'journaliser',AsyncMock()),patch.object(documents_file,'annoncer',AsyncMock()):
      for i in range(3):
       documents_file._maj(cle,prochain=0)
       # Le progrès d'une autre tâche ne doit pas masquer la stagnation.
       dossiers.etape(self.uid,self.fil,'autre','section:'+str(i),{'texte':'autre'})
       with dossiers.base() as c:job=dict(c.execute('SELECT * FROM file_documentaire WHERE id=?',(cle,)).fetchone())
       await documents_file.traiter(job)
     with dossiers.base() as c:job=dict(c.execute('SELECT * FROM file_documentaire WHERE id=?',(cle,)).fetchone())
     self.assertEqual(job['statut'],'bloque');self.assertEqual(job['essais'],3)
     self.assertEqual(json.loads(job['resultat'])['etapes_conservees'],0)
    asyncio.run(test())
 def test_citations_reelles(self):
    documents_dossier._analyse_valide({'faits':[{'fait':'surface','citation':'42,5 m²'}],'limites':[]},'Salle : 42,5 m²')
    with self.assertRaises(ValueError):documents_dossier._analyse_valide({'faits':[{'fait':'surface','citation':'52 m²'}],'limites':[]},'Salle : 42,5 m²')
 def test_image_et_word(self):
    fn=fonctions(BACKEND/'agents/router.py',{'passer_la_main_node'})['passer_la_main_node']
    r=asyncio.run(fn({'attachment_text':'WORD : exigences importantes','vision_analysis':'IMAGE : trois rubriques','attachment_name':'rc.png'}))
    self.assertIn('WORD : exigences importantes',r['attachment_text']);self.assertIn('IMAGE : trois rubriques',r['attachment_text'])
 def test_tableau_apres_couverture_et_autres_feuilles(self):
    from openpyxl import Workbook
    wb=Workbook();wb.active.title='Couverture';wb.active['A1']='Couverture'
    ws=wb.create_sheet('Quantités');ws['B80']='Carrelage';ws['C80']=42.5;ws['D80']='m²';ws['C81']='=C80*2'
    b=io.BytesIO();wb.save(b);texte=lire('dpgf.xlsx',b.getvalue())
    self.assertIn('Feuille Quantités — ligne 80',texte);self.assertIn('42.5',texte);self.assertIn('FORMULE SANS VALEUR CALCULÉE',texte)
 def test_brouillons_isoles_et_cloture(self):
    async def test():
     user=SimpleNamespace(id=self.uid)
     a=await bureau.creer_document({'titre':'Même titre','_fil':'A'},user)
     b=await bureau.creer_document({'titre':'Même titre','_fil':'B'},user)
     self.assertNotEqual(a['document_id'],b['document_id'])
     a2=await bureau.creer_document({'titre':'Même titre','_fil':'A'},user);self.assertEqual(a['document_id'],a2['document_id'])
     jeton=fil_courant.set('B')
     try:
      with self.assertRaises(ValueError):atelier.ajouter(a['document_id'],[{'bloc':'paragraphe','texte':'intrus'}],self.uid)
      with self.assertRaises(ValueError):atelier.terminer(a['document_id'],self.uid)
     finally:fil_courant.reset(jeton)
     atelier.abandonner(a['document_id'],self.uid);atelier.abandonner(b['document_id'],self.uid)
    asyncio.run(test())
 def test_source_pas_preuve_production(self):
    fn=fonctions(BACKEND/'agents/agent1.py',{'_blocs_de','_reference_bloc','_blocs_livrables','_productions_du_tour'},{'_TYPES_LIVRABLE':('fichier','visuel')})
    r={'skill':'nas_ouvrir','ok':True,'resultat_masque':json.dumps({'bloc_ui':{'type':'fichier','url':'/api/documents/source'}})}
    self.assertEqual(len(fn['_blocs_livrables']([r])),1);self.assertEqual(fn['_productions_du_tour']([r]),[])
    r['skill']='terminer_document';self.assertEqual(len(fn['_productions_du_tour']([r])),1)
 def test_collecteur_cellules(self):
    from skills.chiffres_sources import textes
    self.assertIn('42.5','\n'.join(textes({'apercu':[{'Local':'Salle','m²':42.5}]})))
 def test_modele_word_styles_logo_et_contenu_neuf(self):
    from docx import Document
    from docx.shared import Pt
    from PIL import Image
    from bureautique.document_modele import preparer_modele
    from bureautique.rendu import rendre
    from bureautique.modele import normaliser_element
    image=io.BytesIO();Image.new('RGB',(20,20),'blue').save(image,format='PNG');image.seek(0)
    doc=Document();doc.styles['Normal'].font.name='Georgia';doc.styles['Normal'].font.size=Pt(12)
    doc.sections[0].header.paragraphs[0].add_run().add_picture(image)
    doc.add_paragraph('ANCIEN PROJET À SUPPRIMER');b=io.BytesIO();doc.save(b)
    jeton=atelier.ouvrir({'titre':'Rapport neuf','format':'docx'},self.uid,self.fil)
    modele=preparer_modele(jeton,self.uid,b.getvalue());sortie=self.rep.name+'/nouveau.docx'
    rendre({'titre':'Rapport neuf','format':'docx','_modele_docx':modele},[normaliser_element({'bloc':'titre','texte':'Étude complète','niveau':1}),normaliser_element({'bloc':'paragraphe','texte':'EXIGENCE ACTUELLE'})],sortie)
    neuf=Document(sortie);self.assertEqual(neuf.styles['Normal'].font.name,'Georgia');self.assertIn('EXIGENCE ACTUELLE','\n'.join(p.text for p in neuf.paragraphs));self.assertNotIn('ANCIEN PROJET','\n'.join(p.text for p in neuf.paragraphs))
    from docx.oxml.ns import qn
    self.assertEqual(len(list(neuf.sections[0].header._element.iter(qn('w:drawing')))),1)
    atelier.abandonner(jeton,self.uid)
 def test_composition_complete_et_reprise_idempotente(self):
    ids=[dossiers.enregistrer(self.uid,self.fil,'source'+str(i)+'.txt','Exigence vérifiée '+str(i)) for i in range(2)]
    compteur=[]
    async def faux(consigne,data,verifier=None,**options):
     compteur.append(consigne)
     if any(x in consigne for x in ('Rédige intégralement','Vérifie le contenu rédigé','Contrôle final du document')):
      self.assertEqual({x['source'] for x in data['pieces_disponibles']},set(ids))
      self.assertEqual({x['nom'] for x in data['pieces_disponibles']},{'source0.txt','source1.txt'})
     if 'Lis TOUT' in consigne:r={'faits':[{'fait':data['texte_numerote'][0]['texte'],'ligne_debut':1,'ligne_fin':1,'nature':'exigence'}],'limites':[]}
     elif 'Établis le plan' in consigne:r={'titre':'Rapport complet','sections':[{'titre':'Analyse','objectif':'Tout couvrir','sources':ids,'mots_cibles':300}],'sources_ecartees':{},'modele_source':None,'pages_max':None,'remplacements_modele':{'Document confidentiel':'Suggestion interne non demandée'}}
     elif 'Rédige intégralement' in consigne:r={'blocs':[{'bloc':'paragraphe','texte':'Exigence vérifiée 0 et Exigence vérifiée 1 sont traitées avec une présentation complète et concrète de chacune des deux exigences du dossier fourni.'}],'preuves':[a['preuve'] for a in data['preuves']],'reserves':[]}
     elif 'Contrôle final du document' in consigne:
      # Sans les pièces, le contrôle traitait les faits réellement fournis
      # comme des inventions et refusait le livrable à chaque reprise.
      self.assertEqual(data['faits_controles'],2)
      self.assertNotIn('remplacements_modele',data['plan'])
      r={'valide':True,'problemes':[]}
     else:r={'valide':True,'problemes':[]}
     if verifier:verifier(r)
     return r
    async def test():
     user=SimpleNamespace(id=self.uid,email='test@example.invalid');data={'_fil':self.fil,'demande':'Rédige un rapport complet depuis les deux pièces.'}
     with patch.object(documents_dossier,'_json',faux), patch.object(documents_dossier,'_controler_reserves',AsyncMock(side_effect=[TimeoutError('interruption après contrôle interne'),{'problemes':[]}])):
      incomplet=await documents_dossier.composer_immediat(data,user);self.assertFalse(incomplet.get('production_verifiee'),incomplet)
      a=await documents_dossier.composer_immediat({**data,'tache':incomplet['tache']},user);self.assertTrue(a.get('production_verifiee'),a)
      self.assertEqual(sum('Contrôle final du document' in c for c in compteur),1)
      n=len(compteur);b=await documents_dossier.composer_immediat(data,user);self.assertEqual(a['url'],b['url']);self.assertEqual(len(compteur),n)
      self.assertEqual(a['sources_lues'],2);self.assertEqual(a['sections_controlees'],1)
    asyncio.run(test())
 def test_reprise_corrige_section_refusee_sans_repartir_de_zero(self):
    source=dossiers.enregistrer(self.uid,self.fil,'source.txt','Le chantier dure six semaines.')
    compte={'redactions':0,'corrections':0}
    async def faux(consigne,data,verifier=None,**options):
     if 'Lis TOUT' in consigne:r={'faits':[{'fait':'Durée six semaines','citation':'six semaines','nature':'exigence'}],'limites':[]}
     elif 'Établis le plan' in consigne:r={'titre':'Rapport','sections':[{'titre':'Planning','objectif':'Durée','sources':[source],'mots_cibles':100}],'sources_ecartees':{},'modele_source':None,'pages_max':None}
     elif 'Rédige intégralement' in consigne:
      compte['redactions']+=1;r={'blocs':[{'bloc':'paragraphe','texte':'Le planning proposé organise le travail sur six semaines à compter de la mise à disposition du chantier, selon la pièce fournie.'}],'preuves':[source+':1'],'reserves':[]}
     elif 'Corrige la section' in consigne:
      compte['corrections']+=1;self.assertEqual(data['problemes'],['Clarifier le caractère proposé']);r=data['redaction']
     elif 'Vérifie le contenu rédigé' in consigne or 'Vérifie la correction' in consigne:
      r={'valide':compte['corrections']>=2,'problemes':[] if compte['corrections']>=2 else ['Clarifier le caractère proposé']}
     else:r={'valide':True,'problemes':[]}
     if verifier:verifier(r)
     return r
    async def test():
     data={'_fil':self.fil,'demande':'Rédige le rapport'};user=SimpleNamespace(id=self.uid)
     with patch.object(documents_dossier,'_json',faux):
      a=await documents_dossier.composer_immediat(data,user);self.assertFalse(a['production_verifiee'])
      self.assertIsNotNone(dossiers.etape(self.uid,self.fil,a['tache'],'revision:0'))
      b=await documents_dossier.composer_immediat({**data,'tache':a['tache']},user)
     self.assertTrue(b.get('production_verifiee'),b);self.assertEqual(compte,{'redactions':1,'corrections':2})
     self.assertIsNone(dossiers.etape(self.uid,self.fil,a['tache'],'revision:0'))
    asyncio.run(test())
 def test_correction_localise_une_note_presente_seulement_dans_les_reserves(self):
    source=dossiers.enregistrer(self.uid,self.fil,'source.txt','Le chantier dure six semaines.')
    note='Signalement du contrôle : déjà corrigé.';corrections=[]
    async def faux(consigne,data,verifier=None,**options):
     if 'Lis TOUT' in consigne:r={'faits':[{'fait':'Durée six semaines','citation':'six semaines','nature':'exigence'}],'limites':[]}
     elif 'Établis le plan' in consigne:r={'titre':'Rapport','sections':[{'titre':'Planning','objectif':'Durée','sources':[source],'mots_cibles':100}],'sources_ecartees':{},'modele_source':None,'pages_max':None}
     elif 'Rédige intégralement' in consigne:r={'blocs':[{'bloc':'paragraphe','texte':'Le planning proposé organise le travail sur six semaines à compter de la mise à disposition du chantier, selon la pièce fournie.'}],'preuves':[source+':1'],'reserves':[note]}
     elif 'Contrôle final' in consigne:r={'valide':bool(corrections),'problemes':[] if corrections else ['Supprimer la note interne '+note]}
     elif 'Localise les corrections' in consigne:
      self.assertEqual(data['sections'][0]['reserves'],[note]);r={'titre':None,'sections':[0]}
     elif 'Corrige cette section' in consigne:
      corrections.append(1);r={**data['redaction'],'reserves':[]}
     else:r={'valide':True,'problemes':[]}
     if verifier:verifier(r)
     return r
    async def test():
     data={'_fil':self.fil,'demande':'Rédige le rapport'};user=SimpleNamespace(id=self.uid)
     with patch.object(documents_dossier,'_json',faux):
      a=await documents_dossier.composer_immediat(data,user);self.assertFalse(a['production_verifiee'])
      b=await documents_dossier.composer_immediat({**data,'tache':a['tache']},user)
     self.assertTrue(b.get('production_verifiee'),b);self.assertEqual(len(corrections),1)
     self.assertEqual(dossiers.etape(self.uid,self.fil,a['tache'],'section:0')['reserves'],[])
    asyncio.run(test())
 def test_controle_interne_reutilise_seulement_un_succes_inchange(self):
    async def test():
     verifier=AsyncMock(side_effect=[{'valide':False,'problemes':['Erreur']},{'valide':True,'problemes':[]},{'valide':True,'problemes':[]},{'valide':True,'problemes':[]}])
     with patch.object(documents_dossier,'_json',verifier):
      appel=lambda regle,data:documents_dossier._controle_interne(self.uid,self.fil,'tache-test',regle,data)
      a=await appel('Règle',{'texte':'Ancien'});self.assertFalse(a['valide'])
      b=await appel('Règle',{'texte':'Ancien'});self.assertTrue(b['valide'])
      self.assertEqual(await appel('Règle',{'texte':'Ancien'}),b);self.assertEqual(verifier.await_count,2)
      await appel('Règle',{'texte':'Corrigé'});self.assertEqual(verifier.await_count,3)
      await appel('Nouvelle règle',{'texte':'Corrigé'});self.assertEqual(verifier.await_count,4)
    asyncio.run(test())
 def test_acquis_refuses_uniquement_avec_une_reserve_effectivement_presente(self):
    plan={'sections':[{'titre':'Offre'},{'titre':'Candidature'}]}
    reserve='Les références de notre entreprise restent à fournir.'
    sections=[{'blocs':[{'bloc':'paragraphe','texte':'Les références fournies en annexe démontrent notre expérience.'}],'reserves':[]},
              {'blocs':[{'bloc':'paragraphe','texte':'Dossier de candidature.'}],'reserves':[reserve]}]
    async def test():
     appels=[]
     async def faux(consigne,data,verifier=None):
      appels.append(1)
      r={'valide':False,'problemes':[{'section':0,'affirmation':'Les références fournies en annexe','reserve':reserve if len(appels)>1 else 'Réserve inventée','raison':'Ne pas présenter les références comme fournies.'}]}
      if verifier:verifier(r)
      return r
     with patch.object(documents_dossier,'_json',faux):
      with self.assertRaisesRegex(ValueError,'Réserve contradictoire'):await documents_dossier._controler_acquis(self.uid,self.fil,'tache-test',plan,sections)
      r=await documents_dossier._controler_acquis(self.uid,self.fil,'tache-test',plan,sections)
     self.assertEqual(len(appels),2);self.assertEqual(r[0]['section'],0)
    asyncio.run(test())
 def test_acquis_non_prouve_sans_reserve_corrige_avant_livraison(self):
    source=dossiers.enregistrer(self.uid,self.fil,'DPGF.txt','Les quantités de l’offre doivent être vérifiées par le candidat.')
    affirmation='Notre décomposition est vérifiée sans erreur de quantités.'
    compte={'corrections':0}
    async def faux(consigne,data,verifier=None,**options):
     if 'Lis TOUT' in consigne:r={'faits':[{'fait':'Vérification exigée','citation':'Les quantités de l’offre doivent être vérifiées par le candidat.','nature':'exigence'}],'limites':[]}
     elif 'Établis le plan' in consigne:r={'titre':'Offre','sections':[{'titre':'Prix','objectif':'Préciser les vérifications','sources':[source],'mots_cibles':100}],'sources_ecartees':{},'modele_source':None,'pages_max':None}
     elif 'Rédige intégralement' in consigne:r={'blocs':[{'bloc':'paragraphe','texte':affirmation+' Notre équipe intervient conformément aux exigences de vérification du dossier de consultation transmis.'}],'preuves':[source+':1'],'reserves':[]}
     elif 'Vérifie UNIQUEMENT les faits propres' in consigne:
      self.assertEqual(data['faits_du_dossier_complet'][0]['source'],source)
      r={'valide':bool(compte['corrections']),'problemes':[] if compte['corrections'] else [{'section':0,'affirmation':affirmation,'reserve':None,'raison':'Le DPGF exige une vérification mais ne prouve pas sa réalisation.'}]}
     elif 'Corrige cette section' in consigne:
      compte['corrections']+=1
      r={'blocs':[{'bloc':'paragraphe','texte':'La décomposition et les quantités de l’offre restent à vérifier par l’entreprise avant le dépôt ; cette vérification est exigée par le dossier et ne peut pas être annoncée comme déjà réalisée.'}],'preuves':[source+':1'],'reserves':['Vérification des quantités à confirmer.']}
     else:r={'valide':True,'problemes':[]}
     if verifier:verifier(r)
     return r
    async def test():
     data={'_fil':self.fil,'demande':'Rédige le rapport de prix'};user=SimpleNamespace(id=self.uid)
     with patch.object(documents_dossier,'_json',faux):
      a=await documents_dossier.composer_immediat(data,user);self.assertFalse(a['production_verifiee'])
      b=await documents_dossier.composer_immediat({**data,'tache':a['tache']},user)
     self.assertTrue(b.get('production_verifiee'),b);self.assertEqual(compte['corrections'],1)
     from docx import Document
     texte='\n'.join(p.text for p in Document(atelier.chemin_fichier(b['document_id'],self.uid)).paragraphs)
     self.assertNotIn(affirmation,texte);self.assertIn('restent à vérifier',texte)
    asyncio.run(test())
 def test_timeout_documentaire_utilise_secours_sans_retentative_inutile(self):
    from contextlib import asynccontextmanager
    from typing import Any,Optional
    import logging
    @asynccontextmanager
    async def porte():yield
    async def test():
     for documentaire,chain,attendu in [(True,[('test','lent'),('test','secours')],1),(False,[('test','lent'),('test','secours')],1),(True,[('test','lent')],2)]:
      lent=SimpleNamespace(ainvoke=AsyncMock(side_effect=TimeoutError('délai')))
      secours=SimpleNamespace(ainvoke=AsyncMock(return_value=SimpleNamespace(content='{"valide":true}')))
      ns={'asyncio':asyncio,'Any':Any,'Optional':Optional,'LLMTier':object,'logger':logging.getLogger('test.router'),
          'settings':SimpleNamespace(llm_max_retries=2,llm_retry_base_delay=0,ollama_model_light=''),
          '_tier_chain':lambda _:chain,'_filtrer_quarantaine':lambda x:x,
          '_build_model':lambda p,m,*a: lent if m=='lent' else secours,
          '_is_hard_fail':lambda e:False,'_contenu_vide':lambda r:not r.content,
          'tier_timeout':lambda _:10,'tier_max_tokens':lambda _:100}
      tree=ast.parse((BACKEND/'llm/router.py').read_text())
      classe=next(n for n in tree.body if isinstance(n,ast.ClassDef) and n.name=='ResilientLLM')
      exec(compile(ast.Module(body=[classe],type_ignores=[]),'routeur_test','exec'),ns)
      with patch.dict(sys.modules,{'llm.concurrence':SimpleNamespace(porte_llm=porte),'llm.budget':SimpleNamespace(delai_disponible=lambda x:x),'llm.compteur':SimpleNamespace(ajouter=lambda *a:None)}):
       llm=ns['ResilientLLM'](SimpleNamespace(value='complex'))
       if len(chain)==1:
        with self.assertRaises(RuntimeError):await llm.ainvoke([],_secours_timeout_documentaire=documentaire)
       else:
        r=await llm.ainvoke([],_secours_timeout_documentaire=documentaire);self.assertIn('valide',r.content)
        self.assertNotIn('_secours_timeout_documentaire',secours.ainvoke.call_args.kwargs)
      self.assertEqual(lent.ainvoke.await_count,attendu)
    asyncio.run(test())
 def test_correction_reinitialise_le_compteur_sans_effacer_les_preuves(self):
    tache='tache-test';jeton=atelier.ouvrir({'titre':'Brouillon','format':'docx'},self.uid,self.fil)
    dossiers.etape(self.uid,self.fil,tache,'jeton',jeton)
    dossiers.etape(self.uid,self.fil,tache,'suivi_controle',{'cles':['controle_faits:ancien']})
    dossiers.etape(self.uid,self.fil,tache,'controle_faits:ancien',{'problemes':[]})
    dossiers.etape(self.uid,self.fil,tache,'analyse:source',{'faits':['preuve conservée']})
    asyncio.run(documents_dossier._a_corriger(self.uid,self.fil,tache,jeton,{'problemes':['Corriger']}))
    self.assertIsNone(dossiers.etape(self.uid,self.fil,tache,'suivi_controle'))
    self.assertIsNone(dossiers.etape(self.uid,self.fil,tache,'jeton'))
    self.assertEqual(dossiers.etape(self.uid,self.fil,tache,'controle_faits:ancien'),{'problemes':[]})
    self.assertEqual(dossiers.etape(self.uid,self.fil,tache,'analyse:source'),{'faits':['preuve conservée']})
 def test_correction_titre_sans_reecrire_sections(self):
    source=dossiers.enregistrer(self.uid,self.fil,'source.txt','Le chantier dure six semaines.');compte={'controle':0,'redaction':0}
    async def faux(consigne,data,verifier=None,**options):
     if 'Lis TOUT' in consigne:r={'faits':[{'fait':'Durée','citation':'six semaines','nature':'exigence'}],'limites':[]}
     elif 'Établis le plan' in consigne:r={'titre':'Rapport ancien chantier','sections':[{'titre':'Planning','objectif':'Durée','sources':[source],'mots_cibles':100}],'sources_ecartees':{},'modele_source':None,'pages_max':None}
     elif 'Rédige intégralement' in consigne:
      compte['redaction']+=1;r={'blocs':[{'bloc':'paragraphe','texte':'Le planning proposé organise le travail sur six semaines à compter de la mise à disposition du chantier, selon la pièce fournie.'}],'preuves':[source+':1'],'reserves':[]}
     elif 'Localise les corrections' in consigne:r={'titre':'Rapport du chantier','sections':[]}
     elif 'Corrige cette section' in consigne:raise AssertionError('Le corps valide ne doit pas être réécrit pour le titre')
     elif 'Contrôle final' in consigne:
      compte['controle']+=1;r={'valide':compte['controle']>1,'problemes':[] if compte['controle']>1 else ['Titre : référence obsolète à ancien chantier']}
     else:r={'valide':True,'problemes':[]}
     if verifier:verifier(r)
     return r
    async def test():
     data={'_fil':self.fil,'demande':'Rédige un rapport'};user=SimpleNamespace(id=self.uid)
     with patch.object(documents_dossier,'_json',faux):
      a=await documents_dossier.composer_immediat(data,user);self.assertFalse(a['production_verifiee'])
      b=await documents_dossier.composer_immediat({**data,'tache':a['tache']},user)
     self.assertTrue(b.get('production_verifiee'),b);self.assertEqual(compte,{'controle':2,'redaction':1})
     self.assertEqual(dossiers.etape(self.uid,self.fil,a['tache'],'plan')['titre'],'Rapport du chantier')
    asyncio.run(test())
 def test_enveloppe_section_normalisee_sans_relacher_preuves(self):
    contenu={'blocs':[{'bloc':'paragraphe','texte':'Le planning proposé organise le travail sur six semaines à compter de la mise à disposition du chantier, selon la pièce fournie.'}],'preuves':['source:1'],'reserves':[]}
    r={'redaction':contenu.copy()};documents_dossier._section_valide(r,{'source:1'});self.assertEqual(r,contenu)
    with self.assertRaises(ValueError):documents_dossier._section_valide({'redaction':contenu.copy(),'preuves':['autre:1']},{'source:1'})
    with self.assertRaises(ValueError):documents_dossier._section_valide({'redaction':{**contenu,'preuves':['fausse:1']}},{'source:1'})
 def test_reparation_json_recoit_la_reponse_et_erreur_masquees(self):
    router=SimpleNamespace(LLMTier=SimpleNamespace(STANDARD='standard',COMPLEX='complex'),get_llm=None)
    from security.anonymizer import anonymizer
    appels=[]
    invalide=json.dumps({'faits':[{'fait':'Exigence','citation':'[ORG_1] 99 dB'}],'limites':[]})
    valide=json.dumps({'faits':[{'fait':'Exigence','citation':'[ORG_1] 19 dB'}],'limites':[]})
    class Modele:
     async def ainvoke(self,messages,**options):
      appels.append((messages,options));return SimpleNamespace(content=invalide if len(appels)==1 else valide)
    def masquer(chunks,carte):return [x.replace('SOCIETE_TEST_PRIVEE','[ORG_1]') for x in chunks],{'[ORG_1]':'SOCIETE_TEST_PRIVEE'}
    async def test():
     with patch.dict(sys.modules,{'llm.router':router}),patch.object(router,'get_llm',return_value=Modele()),patch.object(anonymizer,'anonymize_chunks',side_effect=masquer),patch.object(anonymizer,'rehydrate',side_effect=lambda t,c:t.replace('[ORG_1]','SOCIETE_TEST_PRIVEE')):
      return await documents_dossier._json('Extrais les faits',{'texte':'SOCIETE_TEST_PRIVEE 19 dB'},lambda r:documents_dossier._analyse_valide(r,'SOCIETE_TEST_PRIVEE 19 dB'),extraction=True)
    r=asyncio.run(test());self.assertEqual(r['faits'][0]['citation'],'SOCIETE_TEST_PRIVEE 19 dB')
    self.assertEqual(len(appels),2);self.assertEqual(appels[1][0][2].content,invalide)
    self.assertIn('99 dB',appels[1][0][3].content)
    self.assertNotIn('SOCIETE_TEST_PRIVEE',' '.join(m.content for m in appels[1][0]))
    self.assertTrue(appels[0][1]['_extraction_documentaire'])
 def test_citations_par_lignes_recopiees_par_le_serveur(self):
    texte='Titre\nPerformance : 19 dB\nPour les circulations.\n'
    r={'faits':[{'fait':'Le seuil est 19 dB pour les circulations','ligne_debut':2,'ligne_fin':3}],'limites':[]}
    documents_dossier._analyse_valide(r,texte)
    self.assertEqual(r['faits'][0]['citation'],'Performance : 19 dB\nPour les circulations.')
    for debut,fin in [(0,2),(2,5),(3,2),(True,2)]:
     with self.assertRaises(ValueError):documents_dossier._analyse_valide({'faits':[{'fait':'Fait','ligne_debut':debut,'ligne_fin':fin}],'limites':[]},texte)
    with self.assertRaises(ValueError):documents_dossier._analyse_valide({'faits':[{'fait':'Fait','ligne_debut':2,'ligne_fin':2,'citation':'Performance : 99 dB'}],'limites':[]},texte)
 def test_passage_long_par_lignes_reste_borne_et_exact(self):
    texte='Clause initiale. '+('Exigence technique du marché. '*90)
    r={'faits':[{'fait':'Une clause longue reste une preuve exacte','ligne_debut':1,'ligne_fin':1}],'limites':[]}
    documents_dossier._analyse_valide(r,texte);self.assertEqual(r['faits'][0]['citation'],texte)
    tableau='Tableau de synthèse.\n'+('Responsable : toutes les entreprises ; action : tri.\n'*150)
    r={'faits':[{'fait':'Le tableau attribue des actions aux entreprises','ligne_debut':1,'ligne_fin':len(tableau.splitlines())}],'limites':[]}
    documents_dossier._analyse_valide(r,tableau);self.assertEqual(r['faits'][0]['citation'],tableau.rstrip('\n'))
    with self.assertRaisesRegex(ValueError,'Fait 1'):
     documents_dossier._analyse_valide({'faits':[{'fait':'Une clause','ligne_debut':1,'ligne_fin':1}],'limites':[]},'x'*(dossiers.TAILLE_FRAGMENT+1))
 def test_rejeu_export_si_disponible(self):
    chemin=os.environ.get('EXPORT_LANGFUSE')
    if not chemin:self.skipTest('Export réel non demandé dans cette recette portable.')
    rows=json.loads(pathlib.Path(chemin).read_text())
    r=next(r for r in rows if r['traceId']=='eea3233b57e935f157366b6250c4ddfd' and r['name']=='classify')
    texte=json.loads(r['input'])['attachment_text'];ids=dossiers.joindre_texte(self.uid,self.fil,texte)
    self.assertEqual(len(ids),7)
    noms=' '.join(s['nom'] for s in dossiers.manifeste(self.uid,self.fil));self.assertIn('NoticeAcoustique',noms);self.assertIn('CCAP',noms)
    self.assertGreater(sum(len(s['contenu']) for s in dossiers.sources(self.uid,self.fil)),600000)

 def test_lecture_paginee_sans_perte(self):
    texte='Début\n'+('Données contrôlées ligne suivante.\n'*1200)+'CIBLE_AU_FOND'
    source=dossiers.enregistrer(self.uid,self.fil,'long.txt',texte)
    async def test():
     args={'source':source,'_fil':self.fil};lus=[]
     while args:
      r=await documents_dossier.lire(args,SimpleNamespace(id=self.uid));lus.append(r['texte'])
      self.assertLess(len(json.dumps(r,ensure_ascii=False)),12000)
      args=r['pour_continuer']['args'] if r['pour_continuer'] else None
      if args:args['_fil']=self.fil
     self.assertEqual(''.join(lus),texte)
     r=await documents_dossier.lire({'source':source,'_fil':self.fil,'recherche':'CIBLE_AU_FOND'},SimpleNamespace(id=self.uid))
     self.assertIn('CIBLE_AU_FOND',r['texte'])
    asyncio.run(test())
 def test_quantites_calcul_unites_et_conflits(self):
    from skills.quantitatifs import _verifier_ligne,dedoublonner
    preuve='Salle : longueur 5 m ; largeur 4 m. Le carrelage est prévu dans la salle.'
    ligne={'lot':'Sol','poste':'Carrelage','niveau':'RDC','local':'Salle','unite':'m²','formule':'c1*c2',
      'operandes':[{'valeur':'5','unite':'m','preuve':'p','citation':'longueur 5 m'},{'valeur':'4','unite':'m','preuve':'p','citation':'largeur 4 m'}],
      'affectation':{'preuve':'p','citation':'Le carrelage est prévu dans la salle.'}}
    r=_verifier_ligne(ligne,{'p':preuve});self.assertEqual(r['quantite'],'20')
    with self.assertRaises(ValueError):_verifier_ligne({**ligne,'unite':'m³'},{'p':preuve})
    mauvaise=json.loads(json.dumps(ligne));mauvaise['operandes'][0]['valeur']='9'
    with self.assertRaises(ValueError):_verifier_ligne(mauvaise,{'p':preuve})
    self.assertEqual(len(dedoublonner([r,r])[0]),1)
    valides,reserves=dedoublonner([r,{**r,'quantite':'21'}]);self.assertEqual(valides,[]);self.assertTrue(reserves)
 def test_quantitatif_excel_reel_et_reprise(self):
    from skills import quantitatifs
    texte='La salle RDC comporte 42.5 m² de carrelage.'
    source=dossiers.enregistrer(self.uid,self.fil,'CCTP.txt',texte)
    ligne={'lot':'Sol','poste':'Carrelage','niveau':'RDC','local':'Salle','unite':'m²','formule':'c1',
      'operandes':[{'valeur':'42.5','unite':'m²','preuve':source+':q1','citation':texte}],
      'affectation':{'preuve':source+':1','citation':texte}}
    async def faux(consigne,data,verifier=None,**options):
     r={'lignes':[ligne],'reserves':[]} if 'CE fragment' in consigne else {'rejeter':[],'reserves':[]}
     if verifier:verifier(r)
     return r
    analyses=[{'source':source,'nom':'CCTP.txt','preuve':source+':1','fragment':1,'faits':[{'fait':texte,'citation':texte,'nature':'exigence'}],'limites':[]}]
    async def test():
     data={'_fil':self.fil,'demande':'Calcule le carrelage au RDC.'};user=SimpleNamespace(id=self.uid,email='test@example.invalid')
     with patch.object(quantitatifs,'_analyses',AsyncMock(return_value=analyses)),patch.object(quantitatifs,'_json',faux):
      r=await quantitatifs.produire_immediat(data,user);self.assertTrue(r['production_verifiee'],r)
      from openpyxl import load_workbook
      wb=load_workbook(atelier.chemin_fichier(r['document_id'],self.uid),read_only=True,data_only=True)
      try:self.assertEqual(wb['Détail']['E2'].value,42.5);self.assertEqual(wb['Synthèse']['E2'].value,42.5)
      finally:wb.close()
      r2=await quantitatifs.produire_immediat(data,user);self.assertEqual(r['url'],r2['url'])
    asyncio.run(test())
 def test_file_persistante_reprise_et_isolation(self):
    from ressources import documents_file as file
    from types import ModuleType
    from contextvars import ContextVar
    source=dossiers.enregistrer(self.uid,self.fil,'source.txt','La source persiste.')
    user=SimpleNamespace(id=self.uid,email='test@example.invalid',role='admin')
    identite=ModuleType('tasks.identity');identite.charger_executant=AsyncMock(return_value=user)
    concurrence=ModuleType('llm.concurrence');concurrence.PERSONNE=ContextVar('test_personne',default=('',0))
    config=ModuleType('config');config.settings=SimpleNamespace(llm_simultanes_fond=2)
    args={'_fil':self.fil,'demande':'Écris un rapport.','sources':[source]}
    r=file.soumettre(self.uid,self.fil,'document',args);r2=file.soumettre(self.uid,self.fil,'document',args)
    self.assertEqual(r['tache_documentaire'],r2['tache_documentaire']);self.assertEqual(file.etats('autre',self.fil),[])
    jeton=atelier.ouvrir({'titre':'Rapport','format':'docx'},self.uid,self.fil);atelier.ajouter(jeton,[{'bloc':'paragraphe','texte':'Rapport complet et contrôlé.'}],self.uid)
    atelier.terminer(jeton,self.uid)
    reussi={'production_verifiee':True,'document_id':jeton,'url':'/api/documents/'+jeton,'bloc_ui':{'type':'fichier','url':'/api/documents/'+jeton}}
    async def test():
     with patch.dict(sys.modules,{'tasks.identity':identite,'llm.concurrence':concurrence,'config':config}),patch.object(file,'annoncer',AsyncMock()) as annonce,patch.object(file,'journaliser',AsyncMock()),patch.object(documents_dossier,'composer_immediat',AsyncMock(side_effect=[TimeoutError(),reussi])) as composer:
      job=file._candidats()[0];await file.traiter(job)
      self.assertEqual(file.etats(self.uid,self.fil)[0]['statut'],'attente');annonce.assert_not_called()
      file._maj(job['id'],prochain=0)
      await file.traiter(file._candidats()[0]);self.assertEqual(file.etats(self.uid,self.fil)[0]['statut'],'termine');annonce.assert_awaited_once()
      await file.traiter(job);self.assertEqual(composer.await_count,2)
      self.assertEqual(file.soumettre(self.uid,self.fil,'document',args)['url'],reussi['url'])
    asyncio.run(test())
 def test_modele_illustration_corps_reelle(self):
    from docx import Document
    from PIL import Image
    from bureautique.illustrations import extraire
    image=io.BytesIO();Image.new('RGB',(80,60),'red').save(image,format='PNG');image.seek(0)
    doc=Document();doc.add_paragraph('Organigramme entreprise');doc.add_picture(image);b=io.BytesIO();doc.save(b)
    images=extraire(b.getvalue());self.assertEqual(len(images),1);self.assertEqual(images[0]['numero'],1)
    self.assertIn('Organigramme',images[0]['contexte']);self.assertGreater(len(images[0]['octets']),50)
 def test_pdf_graphique_avec_texte_ne_perd_pas_les_relations(self):
    import fitz
    pdf=fitz.open();p=pdf.new_page();p.insert_text((30,30),'Planning relatif : phases A et B, graduations M0 M1 M2 M3.')
    for i in range(24):
     x=30+(i%6)*60;y=80+(i//6)*40
     p.draw_rect(fitz.Rect(x,y,x+45,y+20),color=(1,0,0),fill=(1,0,0))
    texte=lire('schema.pdf',pdf.tobytes());pdf.close()
    self.assertIn('=== Page 1 ===\n[LECTURE VISUELLE REQUISE',texte);self.assertIn('graduations M0',texte)
    pdf=fitz.open();p=pdf.new_page();p.insert_text((30,30),'Une lettre avec un petit encadre ne requiert pas la vision.');p.draw_rect(fitz.Rect(30,50,150,90))
    self.assertNotIn('[LECTURE VISUELLE REQUISE',lire('lettre.pdf',pdf.tobytes()));pdf.close()
 def test_fausse_absence_corrigee_avec_preuve_hors_selection_initiale(self):
    a=dossiers.enregistrer(self.uid,self.fil,'planning.txt','La phase dure six semaines.')
    b=dossiers.enregistrer(self.uid,self.fil,'quantites.txt','Surface du local : vingt metres carres.')
    compte={'redactions':0,'corrections':0,'controles':0}
    faux_texte='La méthode proposée décrit la préparation puis la pose selon la durée documentée, avec un suivi des interfaces et des contrôles de réception.'
    async def faux(consigne,data,verifier=None,**options):
     if 'Lis TOUT' in consigne:r={'faits':[{'fait':data['texte_numerote'][0]['texte'],'ligne_debut':1,'ligne_fin':1,'nature':'exigence'}],'limites':[]}
     elif 'Établis le plan' in consigne:r={'titre':'Dossier','sections':[{'titre':'Organisation','objectif':'Préparer','sources':[a],'mots_cibles':200}],'sources_ecartees':{b:'Le plan initial ne retient pas ce tableau'},'modele_source':None,'pages_max':None}
     elif 'Rédige intégralement' in consigne:
      compte['redactions']+=1;r={'blocs':[{'bloc':'paragraphe','texte':faux_texte}],'preuves':[a+':1'],'reserves':['Quantités non fournies.']}
     elif 'Contrôle ciblé des réserves' in consigne:
      compte['controles']+=1;self.assertIn(b+':1',{x['preuve'] for x in data['faits_dossier']})
      r={'problemes':[{'section':0,'reserve':'Quantités non fournies.','raison':'Surface fournie dans le tableau.','sources':[b+':1']}]}
     elif consigne.startswith('Évalue directement'):
      self.assertIn(b+':1',{x['preuve'] for x in data['preuves_completes_concernees']})
      r={'avis':[{'reserve':'Quantités non fournies.','fondee':False,'raison':'Surface fournie dans le tableau.','sources':[b+':1']}]}
     elif 'Corrige cette section' in consigne:
      compte['corrections']+=1;self.assertIn(b+':1',{x['preuve'] for x in data['preuves']})
      r={'blocs':[{'bloc':'paragraphe','texte':faux_texte+' La surface indiquée dans le tableau est de vingt mètres carrés.'}],'preuves':[a+':1',b+':1'],'reserves':[]}
     elif 'Localise les corrections' in consigne:raise AssertionError('La réserve donne déjà la rubrique à corriger')
     else:r={'valide':True,'problemes':[]}
     if verifier:verifier(r)
     return r
    async def test():
     data={'_fil':self.fil,'demande':'Compose le dossier'};user=SimpleNamespace(id=self.uid)
     with patch.object(documents_dossier,'_json',faux):
      r=await documents_dossier.composer_immediat(data,user);self.assertFalse(r['production_verifiee'],r)
      correction=dossiers.etape(self.uid,self.fil,r['tache'],'correction');self.assertEqual(correction['cibles'],[0])
      s=await documents_dossier.composer_immediat({**data,'tache':r['tache']},user);self.assertTrue(s['production_verifiee'],s)
     self.assertEqual(compte,{'redactions':1,'corrections':1,'controles':1})
    asyncio.run(test())
 def test_citation_invalide_reparee_sans_regenerer_faits_acquis(self):
    texte='Surface : vingt metres carres.\n'+('texte secondaire '*450)+'\nDuree : six semaines.'
    sid=dossiers.enregistrer(self.uid,self.fil,'preuve.txt',texte)
    sources=dossiers.sources(self.uid,self.fil,[sid]);compte={'extractions':0,'reparations':0}
    async def faux(consigne,data,verifier=None,**options):
     if 'Lis TOUT' in consigne:
      compte['extractions']+=1;r={'faits':[{'fait':'Surface de vingt metres carres','ligne_debut':1,'ligne_fin':1,'nature':'quantite'},{'fait':'Duree de six semaines','ligne_debut':1,'ligne_fin':4,'nature':'exigence'}],'limites':[]}
     elif 'Répare UNIQUEMENT' in consigne:
      compte['reparations']+=1;self.assertEqual([x['indice'] for x in data['references_a_corriger']],[1])
      if compte['reparations']==1:raise ValueError('Interruption simulée de la réparation')
      r={'corrections':[{'indice':1,'ligne_debut':3,'ligne_fin':3}]}
     else:raise AssertionError(consigne)
     if verifier:verifier(r)
     return r
    async def test():
     with patch.object(documents_dossier,'_json',faux):
      with self.assertRaises(ValueError):await documents_dossier._analyses(self.uid,self.fil,'tache','Rédige le rapport',sources)
      partiel=dossiers.etape(self.uid,self.fil,'tache','analyse_partielle:'+sid+':1');self.assertEqual(len(partiel['faits']),2)
      r=await documents_dossier._analyses(self.uid,self.fil,'tache','Rédige le rapport',sources)
     self.assertEqual(compte,{'extractions':1,'reparations':2});self.assertEqual(r[0]['faits'][0]['citation'],'Surface : vingt metres carres.')
     self.assertEqual(r[0]['faits'][1]['citation'],'Duree : six semaines.');self.assertEqual(r[0]['faits'][1]['fait'],'Duree de six semaines')
     self.assertIsNone(dossiers.etape(self.uid,self.fil,'tache','analyse_partielle:'+sid+':1'))
    asyncio.run(test())
 def test_reparation_citation_interdit_suppression_et_recriture_fait(self):
    async def test():
     for correction in [[],[{'indice':0,'citation':'Duree : six semaines.','fait':'Duree inventee'}],[{'indice':1,'citation':'Duree : six semaines.'}],[{'indice':0,'citation':'huit semaines'}]]:
      async def faux(consigne,data,verifier=None,**options):
       r={'corrections':correction};verifier(r);return r
      r={'faits':[{'fait':'Duree de six semaines','citation':'incorrect'}],'limites':[]}
      with patch.object(documents_dossier,'_json',faux),self.assertRaises(ValueError):await documents_dossier._reparer_citations(r,'Duree : six semaines.')
    asyncio.run(test())
 def test_vision_tronquee_reessayee_puis_secours_structure(self):
    from contextlib import asynccontextmanager
    from langchain_core.messages import HumanMessage,SystemMessage
    import logging,types
    @asynccontextmanager
    async def porte():yield
    module=types.SimpleNamespace(porte_llm=porte)
    appel=fonctions(BACKEND/'agents/agent2.py',{'_appel_vision'},{'HumanMessage':HumanMessage,'SystemMessage':SystemMessage,'logger':logging.getLogger('test')})['_appel_vision']
    calls=[]
    class Modele:
     def __init__(self,nom):self.nom=nom
     async def ainvoke(self,messages,**kwargs):
      calls.append((self.nom,messages,kwargs))
      return SimpleNamespace(content='texte interrompu' if self.nom=='coupe' else [{'type':'text','text':'{"observations":[]}'}],response_metadata={'finish_reason':'length' if self.nom=='coupe' else 'stop'},usage_metadata={})
    def verifier(t):self.assertIn('observations',json.loads(t))
    async def test():
     with patch.dict(sys.modules,{'llm.concurrence':module}):
      r=await appel([(Modele('coupe'),'coupe'),(Modele('bon'),'bon')],'contexte',[],'piece',consigne_systeme='Lecture seule, JSON',verifier=verifier)
     self.assertEqual(r['model_used'],'bon');self.assertEqual(len(calls),3)
     self.assertEqual(calls[0][2]['max_tokens'],8192);self.assertEqual(calls[1][2]['max_tokens'],16384);self.assertIsInstance(calls[-1][1][0],SystemMessage)
    asyncio.run(test())
 def test_vision_non_structuree_ne_devient_pas_une_preuve(self):
    from contextlib import asynccontextmanager
    from langchain_core.messages import HumanMessage,SystemMessage
    import logging,types
    @asynccontextmanager
    async def porte():yield
    appel=fonctions(BACKEND/'agents/agent2.py',{'_appel_vision'},{'HumanMessage':HumanMessage,'SystemMessage':SystemMessage,'logger':logging.getLogger('test')})['_appel_vision']
    class Modele:
     async def ainvoke(self,*a,**k):return SimpleNamespace(content='Je vous ai prépare un mémoire',response_metadata={},usage_metadata={})
    async def test():
     with patch.dict(sys.modules,{'llm.concurrence':types.SimpleNamespace(porte_llm=porte)}):
      r=await appel([(Modele(),'mauvais')],'contexte',[],'piece',verifier=json.loads)
     self.assertNotIn('analyse',r);self.assertIn('erreur',r)
    asyncio.run(test())
 def test_suspendre_par_identifiant_moteur_reste_dans_le_fil(self):
    from ressources import documents_file
    source=dossiers.enregistrer(self.uid,self.fil,'piece.txt','Une preuve')
    r=documents_file.soumettre(self.uid,self.fil,'document',{'sources':[source],'demande':'Rapport'})
    ident=r['tache_documentaire']
    with dossiers.base() as c:
     d=json.loads(c.execute('SELECT donnees FROM file_documentaire WHERE id=?',(ident,)).fetchone()[0]);d['tache']='identifiant-moteur'
     c.execute('UPDATE file_documentaire SET donnees=? WHERE id=?',(json.dumps(d),ident))
    with self.assertRaises(ValueError):documents_file.piloter(self.uid,'autre-fil','identifiant-moteur')
    with self.assertRaises(ValueError):documents_file.piloter('autre',self.fil,'identifiant-moteur')
    self.assertEqual(documents_file.piloter(self.uid,self.fil,'identifiant-moteur')['statut'],'suspendu')
    self.assertEqual(documents_file.piloter(self.uid,self.fil,ident,reprendre=True)['statut'],'attente')
 def test_illustration_decorative_ne_devient_pas_fausse_preuve(self):
    import types
    from docx import Document
    from PIL import Image
    from bureautique import illustrations
    image=io.BytesIO();Image.new('RGB',(80,60),'blue').save(image,format='PNG');image.seek(0)
    doc=Document();doc.add_paragraph('Cadre de réponse');doc.add_picture(image);b=io.BytesIO();doc.save(b)
    async def vision(*args,**kwargs):
     self.assertIn('AUTRES pièces',kwargs['consigne_systeme']);texte=json.dumps({'type':'decorative','description':'Motif bleu','faits_lisibles':[],'incertitudes':[]});kwargs['verifier'](texte);return {'analyse':texte}
    async def test():
     with patch.dict(sys.modules,{'agents.agent2':types.SimpleNamespace(_appel_vision=vision),'llm.router':types.SimpleNamespace(get_vision_candidates=lambda:[])}),patch('mail.attaches.resoudre',AsyncMock(return_value=([{'octets':b.getvalue()}],[]))):
      r=await illustrations.analyser(self.uid,self.fil,'tache',[{'id':'s','nom':'cadre.docx','reference':'/api/documents/test'}],SimpleNamespace(id=self.uid),'Rédige un document avec toutes les pièces')
     self.assertEqual(r[0]['faits'],[]);self.assertEqual(r[0]['illustration']['type'],'decorative')
     self.assertIsNotNone(dossiers.etape(self.uid,self.fil,'tache','illustration:v2:s:1'))
    asyncio.run(test())
 def test_illustration_informative_conserve_faits_et_refuse_fausse_decoration(self):
    from bureautique.illustrations import _lecture
    r=_lecture(json.dumps({'type':'information','description':'Organigramme','faits_lisibles':['Équipe de deux compagnons'],'incertitudes':['Un nom est illisible']}));self.assertEqual(r['faits_lisibles'],['Équipe de deux compagnons'])
    with self.assertRaises(ValueError):_lecture('Je vais rédiger le mémoire et les autres pièces ne sont pas disponibles.')
    with self.assertRaises(ValueError):_lecture(json.dumps({'type':'decorative','description':'Icône','faits_lisibles':['L’entreprise a deux salariés'],'incertitudes':[]}))
 def test_pdf_hybride_signale_page_non_lue(self):
    import fitz
    pdf=fitz.open();page=pdf.new_page();page.insert_text((60,60),'Page textuelle exploitable avec les contraintes du projet.');pdf.new_page()
    texte=lire('mixte.pdf',pdf.tobytes());pdf.close()
    self.assertIn('=== Page 2 ===\n[LECTURE VISUELLE REQUISE',texte)
 def test_pagination_reelle_si_libreoffice(self):
    import shutil
    if not (shutil.which('libreoffice') or shutil.which('soffice')):self.skipTest('Pagination réelle exécutée dans le conteneur LibreOffice.')
    from docx import Document
    from bureautique.document_modele import verifier_pages,convertir_pdf
    import fitz
    doc=Document();doc.add_paragraph('Première page avec texte réel');doc.add_page_break();doc.add_paragraph('Deuxième page avec texte réel')
    p=self.rep.name+'/pagination.docx';doc.save(p)
    r=verifier_pages(p,1);self.assertFalse(r['conforme']);self.assertEqual(r['pages'],2)
    r=verifier_pages(p,2);self.assertTrue(r['conforme'])
    with fitz.open(stream=convertir_pdf(p),filetype='pdf') as pdf:self.assertEqual(len(pdf),2)

 def test_suspension_et_reprise_conservent_sources(self):
    from ressources import documents_file as file
    source=dossiers.enregistrer(self.uid,self.fil,'piece.txt','Document de référence conservé.')
    r=file.soumettre(self.uid,self.fil,'document',{'demande':'Rédige','sources':[source]});cle=r['tache_documentaire']
    with self.assertRaises(ValueError):file.piloter('intrus',self.fil,cle)
    file.piloter(self.uid,self.fil,cle)
    self.assertEqual(file._candidats(),[])
    jeton=file._COURANTE.set(cle)
    try:
     with self.assertRaises(ValueError):file.verifier_poursuite()
    finally:file._COURANTE.reset(jeton)
    file.piloter(self.uid,self.fil,cle,True);self.assertEqual(len(file._candidats()),1)
    self.assertEqual(len(dossiers.sources(self.uid,self.fil,[source])),1)
 def test_rejet_du_rendu_permet_une_revision(self):
    jeton=atelier.ouvrir({'titre':'À reprendre','format':'docx'},self.uid,self.fil)
    tache='redaction-a-reprendre';dossiers.etape(self.uid,self.fil,tache,'jeton',jeton)
    dossiers.etape(self.uid,self.fil,tache,'section:0',{'blocs':[{'bloc':'paragraphe','texte':'Rédaction sauvegardée.'}]})
    asyncio.run(documents_dossier._a_corriger(self.uid,self.fil,tache,jeton,{'facteur_longueur':.75}))
    self.assertIsNone(dossiers.etape(self.uid,self.fil,tache,'jeton'))
    self.assertTrue(dossiers.etape(self.uid,self.fil,tache,'section:0'))
    self.assertEqual(dossiers.etape(self.uid,self.fil,tache,'correction')['facteur_longueur'],.75)
 def test_pdf_compose_si_libreoffice(self):
    import shutil
    if not shutil.which('libreoffice'):self.skipTest('Conversion du document composé vérifiée dans Docker.')
    source=dossiers.enregistrer(self.uid,self.fil,'source.txt','Exigence contractuelle actuelle.')
    plan={'titre':'PDF contrôlé','sections':[{'titre':'Analyse','sources':[source]}],'pages_max':3,'modele_source':None}
    sections=[{'blocs':[{'bloc':'paragraphe','texte':'Une rédaction nouvelle, concrète et vérifiée, qui reprend les exigences du dossier fourni.'}],'reserves':[]}]
    async def test():
     with patch.object(documents_dossier,'_json',AsyncMock(return_value={'valide':True,'problemes':[]})):
      r=await documents_dossier._rendre(self.uid,self.fil,'pdf-test',{'titre':'PDF contrôlé','demande':'Rédige ce rapport en PDF','format':'pdf'},plan,dossiers.sources(self.uid,self.fil,[source]),sections,SimpleNamespace(id=self.uid))
      self.assertEqual(r['format'],'pdf');self.assertTrue(r['production_verifiee'])
      import fitz
      with fitz.open(atelier.chemin_fichier(r['document_id'],self.uid)) as pdf:self.assertIn('Analyse',''.join(p.get_text() for p in pdf))
    asyncio.run(test())

 def test_modele_plusieurs_sections_conserve_entete_herite(self):
    from docx import Document
    from docx.enum.section import WD_SECTION
    from bureautique.document_modele import preparer_modele
    doc=Document();doc.sections[0].header.paragraphs[0].text='Entreprise — en-tête commun'
    doc.sections[0].footer.paragraphs[0].text='Pied de page commun'
    doc.add_paragraph('Ancien projet première section');doc.add_section(WD_SECTION.NEW_PAGE);doc.add_paragraph('Ancien projet deuxième section')
    self.assertTrue(doc.sections[-1].header.is_linked_to_previous)
    b=io.BytesIO();doc.save(b);jeton=atelier.ouvrir({'titre':'Rapport neuf'},self.uid,self.fil)
    p=preparer_modele(jeton,self.uid,b.getvalue());neuf=Document(p)
    self.assertEqual(len(neuf.sections),1)
    self.assertIn('Entreprise',neuf.sections[0].header.paragraphs[0].text)
    self.assertIn('Pied de page commun',neuf.sections[0].footer.paragraphs[0].text)
    self.assertFalse(any('Ancien projet' in p.text for p in neuf.paragraphs))

 def test_conversion_pdf_en_panne_reprend_word_modele_termine(self):
    from docx import Document
    from types import ModuleType
    from bureautique import document_modele
    import fitz,hashlib
    original=Document();original.sections[0].header.paragraphs[0].text='Entreprise test';original.add_paragraph('Ancien projet')
    b=io.BytesIO();original.save(b);octets=b.getvalue()
    source=dossiers.enregistrer(self.uid,self.fil,'Modele.docx','Informations de présentation.','/api/documents/modele-test',hashlib.sha256(octets).hexdigest())
    plan={'titre':'Rapport PDF','sections':[{'titre':'Analyse','sources':[source]}],'pages_max':None,'modele_source':source}
    sections=[{'blocs':[{'bloc':'paragraphe','texte':'Rédaction nouvelle et contrôlée.'}],'reserves':[]}]
    attaches=ModuleType('mail.attaches');attaches.resoudre=AsyncMock(return_value=([{'octets':octets,'nom':'Modele.docx'}],[]))
    pdf=fitz.open();pdf.new_page();pdfbytes=pdf.tobytes();pdf.close()
    async def test():
     args=(self.uid,self.fil,'reprise-pdf',{'titre':'Rapport PDF','demande':'Rédige','format':'pdf'},plan,dossiers.sources(self.uid,self.fil,[source]),sections,SimpleNamespace(id=self.uid,email='test@example.invalid'))
     with patch.dict(sys.modules,{'mail.attaches':attaches}),patch.object(documents_dossier,'_json',AsyncMock(return_value={'valide':True})),patch.object(document_modele,'convertir_pdf',side_effect=[ValueError('Interruption LibreOffice simulée'),pdfbytes]) as convertir:
      with self.assertRaises(ValueError):await documents_dossier._rendre(*args)
      jeton=dossiers.etape(self.uid,self.fil,'reprise-pdf','jeton');self.assertTrue(atelier.fiche(jeton,self.uid)['fini'])
      r=await documents_dossier._rendre(*args);self.assertTrue(r['production_verifiee'])
      r2=await documents_dossier._rendre(*args);self.assertEqual(r['url'],r2['url']);self.assertEqual(convertir.call_count,2)
    asyncio.run(test())

if __name__=='__main__':unittest.main(verbosity=2)
