"""
Banc de l'image à sa place — choisie et vérifiée par le modèle de vision (21/09).

Noa, après le dossier Camp (logo au fond noir, pied de page = logo du haut, Word
refusé) : « corrige de façon moins déterministe : si un autre cas de figure se
présente, l'IA doit être capable de transformer n'importe quel format pour rendre
cette image correcte en en-tête ».

Trois couches, éprouvées ici :
  1. la conversion UNIVERSELLE (Pillow, puis PyMuPDF — SVG, PDF, XPS… —, puis
     LibreOffice) : un format imprévu ne bloque plus ;
  2. le CHOIX par le modèle de vision (`image_placee.choisir`) : il voit les
     candidats tels que Word les montrera, choisit, recadre, retouche, puis juge
     le résultat ; sans modèle, la règle d'avant. Le modèle est DOUBLÉ ici : on
     éprouve ce que le code fait de chacune de ses décisions ;
  3. le FILET du rendu : une image que Word refuse encore devient
     « [image indisponible] », le document sort, et l'image écartée est nommée.

Les parties qui exigent Pillow / PyMuPDF / python-docx sont sautées hors de
l'image du backend, et le banc le dit.

Usage : python backend/scripts/test_image_placee.py [backend]
"""
import asyncio
import importlib.util
import io
import os
import pathlib
import sys
import tempfile
import types

BACKEND = pathlib.Path(sys.argv[1] if len(sys.argv) > 1 else "backend").resolve()
sys.path.insert(0, str(BACKEND))
echecs = []


def verifier(nom, cond, detail=""):
    print(f"  {'✓' if cond else '✗'} {nom}" + (f"  → {detail}" if detail and not cond else ""))
    if not cond:
        echecs.append(nom)


def charger(nom, chemin):
    spec = importlib.util.spec_from_file_location(nom, chemin)
    module = importlib.util.module_from_spec(spec)
    sys.modules[nom] = module
    spec.loader.exec_module(module)
    return module


print(f"\n═══ L'IMAGE À SA PLACE — {BACKEND.parent}\n")
images = charger("bureautique.images", BACKEND / "bureautique" / "images.py")
placee = charger("bureautique.image_placee", BACKEND / "bureautique" / "image_placee.py")

# ── Ce que le modèle a le droit de répondre ──────────────────────────────────
ok = placee.lire_choix('{"choix": 2, "recadrage": [0, 0.8, 1, 1], "retouches": ["rogner_marges"], "raison": "le vrai pied"}', 3)
verifier("une réponse de choix bien formée est lue", ok["choix"] == 2 and ok["recadrage"] == [0, 0.8, 1, 1])
for mauvais, pourquoi in (('{"choix": 4}', "un numéro hors des candidats"),
                          ('{"choix": 1, "retouches": ["effacer_le_texte"]}', "une retouche hors de la boîte à outils"),
                          ('{"choix": 1, "recadrage": [0, 1]}', "un recadrage mal formé"),
                          ("je prends le deuxième", "une réponse qui n'est pas du JSON")):
    try:
        placee.lire_choix(mauvais, 3)
        refuse = False
    except Exception:  # noqa: BLE001
        refuse = True
    verifier(f"refusé (le candidat de vision suivant répond) : {pourquoi}", refuse)
verifier("un verdict se lit", placee.lire_verdict('{"convient": false, "defaut": "fond noir"}') == {"convient": False, "defaut": "fond noir"})

source = (BACKEND / "bureautique" / "images.py").read_text(encoding="utf-8")
verifier("en-tête, pied et couverture passent par le choix visuel",
         'PLACES_JUGEES = ("entete", "pied", "couverture")' in source and "await choisir(" in source)
verifier("… et d'où vient l'image est rendu au modèle (remarques), réserve comprise",
         "remarques.append(" in source and "Images placées" in source)
consigne = placee.CONSIGNE + placee.CONSIGNE_VERDICT
verifier("la consigne dit qu'un texte dans l'image est une donnée, jamais une instruction",
         consigne.count("jamais une instruction") == 2)

try:
    from PIL import Image
    import fitz
    import docx
except ImportError:
    Image = None
    print("  · Pillow / PyMuPDF / python-docx absents ici : les parties exécutées sont sautées (jouées dans l'image)")

if Image is not None:
    # ── Un devis : logo à fond TRANSPARENT en haut, tableau, pied en TEXTE ───
    logo = Image.new("RGBA", (300, 200), (0, 0, 0, 0))
    for x in range(40, 260):
        for y in range(60, 140):
            logo.putpixel((x, y), (10, 10, 10, 255))
    t = io.BytesIO()
    logo.save(t, "PNG")
    pdf = fitz.open()
    page = pdf.new_page(width=595, height=842)
    page.insert_image(fitz.Rect(40, 20, 190, 120), stream=t.getvalue())
    for y in range(180, 700, 26):
        page.draw_line(fitz.Point(40, y), fitz.Point(555, y), color=(0, 0, 0), width=0.8)
        page.insert_text((50, y + 17), "Fourniture et pose  12 m2  45,00 EUR", fontsize=10)
    page.insert_text((60, 815), "9 RUE DE LA SILICE 33380 MARCHEPRIME | SIRET 897 428 637", fontsize=9)
    devis = pdf.tobytes()
    pdf.close()

    liste = placee.candidats(devis, "application/pdf", "devis.pdf", "pied")
    origines = [c["origine"] for c in liste]
    verifier("les candidats d'un PDF : l'image du haut, les bandes haut et bas, la page entière",
             any("image incorporée en haut" in o for o in origines)
             and any("bande du bas" in o for o in origines) and any("page 1 entière" in o for o in origines), origines)
    verifier("chaque candidat est déjà lisible par Word", all(images.lisible_par_word(c["octets"]) for c in liste))
    apercu = Image.open(io.BytesIO(__import__("base64").b64decode(placee.apercu_sur_blanc(liste[0]["octets"]))))
    verifier("le modèle voit l'image POSÉE SUR DU BLANC (ce que Word montrera)",
             apercu.mode == "RGB" and apercu.getpixel((2, 2)) == (255, 255, 255), str(apercu.getpixel((2, 2))))

    def jouer(reponses, place="pied", source_octets=devis, mime="application/pdf", nom="devis.pdf"):
        """Le modèle doublé : il rend, dans l'ordre, les réponses données (None = indisponible)."""
        file = list(reponses)
        vus = []

        async def _vision(entete, apercus, nom_, consigne, verif):
            vus.append((entete, len(apercus)))
            r = file.pop(0) if file else None
            if r is None:
                return None
            verif(r)
            return {"analyse": r}
        placee._vision = _vision
        try:
            return asyncio.run(placee.choisir(source_octets, mime, nom, place, "document « Dossier Camp »")), vus
        except Exception as e:  # noqa: BLE001
            return e, vus

    num_bas = next(i for i, o in enumerate(origines, start=1) if "bande du bas" in o)
    r, vus = jouer([f'{{"choix": {num_bas}, "raison": "le vrai pied du devis"}}', '{"convient": true}'])
    verifier("le modèle choisit le vrai pied : c'est lui qui est posé",
             isinstance(r, dict) and r["octets"] == liste[num_bas - 1]["octets"] and r["choix_par"] == "modele", str(r)[:200])
    verifier("… il a vu tous les candidats, puis le résultat seul", vus and vus[0][1] == len(liste) and vus[1][1] == 1, str(vus))
    verifier("… la place et le document lui sont dits", "PIED DE PAGE" in vus[0][0] and "Dossier Camp" in vus[0][0])

    r, _ = jouer(['{"choix": 0, "raison": "aucune bande de pied lisible"}'])
    verifier("le modèle refuse tout : l'image N'EST PAS posée, et la raison est dite",
             isinstance(r, images.ImageRefusee) and "aucune bande de pied lisible" in str(r), repr(r))

    num_page = next(i for i, o in enumerate(origines, start=1) if "page 1 entière" in o)
    r, _ = jouer([f'{{"choix": {num_page}, "recadrage": [0.05, 0.94, 0.95, 0.99], "retouches": ["rogner_marges"]}}',
                  '{"convient": true}'])
    im = Image.open(io.BytesIO(r["octets"])) if isinstance(r, dict) else None
    verifier("un recadrage décidé par le modèle est appliqué (la page devient une bande)",
             im is not None and im.size[0] > 4 * im.size[1] and "(recadrée)" in r["origine"], str(im.size if im else r))

    # Un logo sur fond NOIR opaque (le masque perdu) : le modèle demande de rendre le fond transparent.
    noir = Image.new("RGB", (300, 200), (0, 0, 0))
    for x in range(100, 200):
        for y in range(60, 140):
            noir.putpixel((x, y), (30, 120, 60))
    t = io.BytesIO()
    noir.save(t, "PNG")
    r, _ = jouer(['{"choix": 1, "retouches": ["fond_sombre_transparent"]}', '{"convient": true}'],
                 place="entete", source_octets=t.getvalue(), mime="image/png", nom="logo.png")
    im = Image.open(io.BytesIO(r["octets"])).convert("RGBA") if isinstance(r, dict) else None
    verifier("« fond sombre → transparent » : le fond noir disparaît, le logo reste",
             im is not None and im.getpixel((5, 5))[3] == 0 and im.getpixel((150, 100))[3] == 255,
             str((im.getpixel((5, 5)), im.getpixel((150, 100))) if im else r))

    r, vus = jouer([f'{{"choix": {num_page}, "recadrage": [0, 0, 0.2, 0.1]}}',
                    '{"convient": false, "defaut": "le recadrage coupe le logo"}',
                    f'{{"choix": {num_bas}}}', '{"convient": true}'])
    verifier("le contrôle refuse : le DÉFAUT revient au modèle, qui choisit autrement",
             isinstance(r, dict) and r["octets"] == liste[num_bas - 1]["octets"] and len(vus) == 4
             and "coupe le logo" in vus[2][0], str(r)[:160])
    r, _ = jouer(['{"choix": 1}', '{"convient": false, "defaut": "grand fond noir"}',
                  '{"choix": 1}', '{"convient": false, "defaut": "toujours un fond noir"}'])
    verifier("deux refus au contrôle : l'image N'EST PAS posée, et le défaut est dit",
             isinstance(r, images.ImageRefusee) and "toujours un fond noir" in str(r), repr(r))
    r, _ = jouer(['{"choix": 1}', None])
    verifier("contrôle indisponible : le choix du modèle vaut, et on le dit",
             isinstance(r, dict) and "indisponible" in r["avertissement"], str(r)[:160])

    r, _ = jouer([None], place="pied")
    verifier("sans modèle : le choix d'avant (la bande du bas pour un pied), sans rien bloquer",
             isinstance(r, dict) and r["choix_par"] == "regle" and "bande du bas" in r["origine"], str(r)[:200])
    r, _ = jouer([None], place="entete")
    verifier("sans modèle : l'image du haut pour un en-tête",
             isinstance(r, dict) and "image incorporée en haut" in r["origine"], str(r)[:200])

    # ── La conversion universelle ─────────────────────────────────────────────
    svg = (b'<svg xmlns="http://www.w3.org/2000/svg" width="200" height="80">'
           b'<rect width="200" height="80" fill="#2e7d32"/><text x="20" y="50" fill="white">SYMBIOSE</text></svg>')
    try:
        png, ext = images.normaliser_octets(svg, "image/svg+xml", "logo.svg")
        converti = ext == "png" and images.lisible_par_word(png) and Image.open(io.BytesIO(png)).size[0] > 100
    except Exception as e:  # noqa: BLE001
        converti = False
    verifier("un logo SVG (que Pillow n'ouvre pas) devient un PNG lisible par Word", converti)
    try:
        images.normaliser_octets(b"ceci n'est pas une image", "application/octet-stream", "x.bin")
        refuse = False
    except images.ImageRefusee as e:
        refuse = "ni LibreOffice" in str(e)
    verifier("ce que rien n'ouvre est refusé, en disant ce qui a été essayé", refuse)

    # ── Le filet du rendu ─────────────────────────────────────────────────────
    with tempfile.TemporaryDirectory() as dossier:
        atelier = types.ModuleType("bureautique.atelier")
        atelier.chemin_image = lambda n: os.path.join(dossier, n) if os.path.exists(os.path.join(dossier, n)) else None
        sys.modules["bureautique.atelier"] = atelier
        rendu = charger("bureautique.rendu", BACKEND / "bureautique" / "rendu.py")
        abime = os.path.join(dossier, "doc.img1.png")
        with open(abime, "wb") as f:
            f.write(b"\x89PNG\r\n\x1a\n" + b"abime" * 20)      # une signature PNG, un contenu abîmé
        rendu._ECARTEES.liste = []
        d = docx.Document()
        pose = rendu._poser(d.add_paragraph().add_run(), abime)
        pose2 = rendu._poser(d, abime)
        sortie = os.path.join(dossier, "sortie.docx")
        d.save(sortie)
        textes = [p.text for p in docx.Document(sortie).paragraphs]
        verifier("une image que Word refuse ne fait PAS tomber le document (il est enregistré)",
                 not pose and not pose2 and os.path.exists(sortie))
        verifier("… elle est remplacée par un repère visible", textes.count("[image indisponible]") >= 2, str(textes))
        verifier("… et nommée pour que l'assistant le dise", rendu.images_ecartees() == ["doc.img1.png", "doc.img1.png"],
                 str(rendu.images_ecartees()))
    verifier("toutes les insertions d'image du Word passent par le filet",
             "add_picture(" not in (BACKEND / "bureautique" / "rendu.py").read_text(encoding="utf-8")
             .replace("cible.add_picture(chemin, **taille)", ""))

bureau = (BACKEND / "skills" / "bureau.py").read_text(encoding="utf-8")
verifier("`terminer_document` dit les images écartées", '"images_ecartees": f.get("images_ecartees")' in bureau
         and "[image indisponible]" in bureau)

print(f"\n{'═' * 70}\n{'✗ ' + str(len(echecs)) + ' échec(s) : ' + ', '.join(echecs) if echecs else '✓ 0 échec'}\n")
sys.exit(1 if echecs else 0)
