"""
Banc du WORD SOIGNÉ — la mise en page de la maison, par défaut (15/09).

Relevé de Noa (Duret, mémoire technique) : « la mise en page est horrible, les
phrases se chevauchent et il n'y a que des bullet points ; possible de faire un
skill pour qu'il fasse des beaux Word ? ». Le rendu partait du gabarit VIERGE de
python-docx : polices de THÈME (résolues au hasard selon le logiciel — c'est là
que les lignes se tassent), aucun interligne, ni page de garde, ni sommaire,
tableaux nus ; et une énumération écrite dans un paragraphe restait une ligne.

CE QUE CE BANC PROUVE (le Word est PRODUIT puis ROUVERT et son XML lu) :
  · styles explicites : police nommée partout (aucun attribut de thème),
    interligne 1,15, titres à la couleur de la charte gardés avec la suite ;
  · document long : page de garde (titre, filet de la charte, date), puis
    sommaire (champ TOC) et mise à jour des champs à l'ouverture — chacun À SA
    PLACE dans l'ordre du schéma (sinon Word dit « contenu illisible ») ;
  · document court : ni page de garde ni sommaire ; `page_de_garde: false` les
    retire ;
  · un paragraphe à tirets devient une vraie liste, « **gras** » du gras ;
  · tableau : en-tête à la couleur de la charte, répété à chaque page, toute la
    largeur utile ;
  · le modèle garde les retours à la ligne d'un paragraphe ; le PDF les rend.
Exige python-docx (sinon le banc le DIT et saute les rendus).

Usage : python backend/scripts/test_word_soigne.py [backend]
"""
import os
import pathlib
import sys
import tempfile
import types
import zipfile

racine = pathlib.Path(sys.argv[1] if len(sys.argv) > 1 else "backend").resolve()
sys.path.insert(0, str(racine))
os.environ.setdefault("DOCUMENTS_DIR", tempfile.mkdtemp(prefix="banc-word-"))
echecs = []


def verifier(nom, cond, detail=""):
    print(f"  {'✓' if cond else '✗'} {nom}" + (f"  → {str(detail)[:300]}" if detail and not cond else ""))
    if not cond:
        echecs.append(nom)


sys.modules.setdefault("emails", types.ModuleType("emails"))
sys.modules["emails.marque"] = types.SimpleNamespace(MARQUE={"couleur": "#0A6FB4", "fond": "#0B3A5B"})
from bureautique import modele  # noqa: E402

print("1. Le modèle garde ce qu'il faut")
p = modele.normaliser_element({"bloc": "paragraphe", "texte": "Pièces :\n- DC1\n  - Kbis  \n\n"})
verifier("les retours à la ligne d'un paragraphe restent", p and p["texte"] == "Pièces :\n- DC1\n- Kbis", p)
t = modele.normaliser_element({"bloc": "titre", "texte": "Un\ntitre"})
verifier("un titre, lui, reste sur une ligne", t and t["texte"] == "Un titre", t)
e = modele.normaliser_entete({"titre": "X", "page_de_garde": False})
verifier("`page_de_garde` et `sommaire` passent (absents = décidés par le rendu)",
         e["page_de_garde"] is False and e["sommaire"] is None, e)

try:
    import docx  # noqa: F401
except ImportError:
    print("\n  (python-docx absent : rendus sautés — jouer ce banc dans le conteneur ou un venv)")
    docx = None

if docx is not None:
    from bureautique import rendu
    from docx import Document
    from docx.oxml.ns import qn

    def rendre(entete, blocs):
        chemin = tempfile.mktemp(suffix=".docx")
        rendu._docx(modele.normaliser_entete(entete),
                    [x for x in (modele.normaliser_element(b) for b in blocs) if x], chemin)
        return chemin

    long = ([{"bloc": "titre", "texte": f"Partie {i}", "niveau": 1} for i in range(1, 7)]
            + [{"bloc": "paragraphe", "texte": "Notre méthode repose sur une préparation soignée des supports."}] * 6
            + [{"bloc": "paragraphe", "texte": "Pièces à fournir :\n- **Lettre de candidature** (DC1)\n- Attestation d'assurance"},
               {"bloc": "tableau", "entetes": ["Moyen", "Quantité"], "lignes": [["Ponceuse", "2"], ["Laser", "1"]]}])
    chemin = rendre({"titre": "Mémoire technique", "sous_titre": "Lots 11 et 12", "entete": "DURET & SOLS",
                     "pied": "Affaire 2026-71"}, long)
    d = Document(chemin)
    with zipfile.ZipFile(chemin) as z:
        xml = z.read("word/document.xml").decode("utf-8")
        reglages = z.read("word/settings.xml").decode("utf-8")
        styles_xml = z.read("word/styles.xml").decode("utf-8")

    print("2. Les styles de la maison")
    normal = d.styles["Normal"]
    polices = normal.element.rPr.find(qn("w:rFonts"))
    verifier("police nommée pour tous les jeux de caractères, sans attribut de thème",
             polices is not None and polices.get(qn("w:ascii")) == rendu.POLICE
             and polices.get(qn("w:eastAsia")) == rendu.POLICE and polices.get(qn("w:asciiTheme")) is None)
    verifier("interligne 1,15 et espace après les paragraphes",
             abs((normal.paragraph_format.line_spacing or 0) - 1.15) < 0.01
             and normal.paragraph_format.space_after is not None)
    h1 = d.styles["Heading 1"]
    verifier("les titres sont à la couleur de la charte et restent avec la suite",
             str(h1.font.color.rgb) == modele.COULEURS["charte"] and h1.paragraph_format.keep_with_next)
    verifier("aucun style de titre ne garde une police de thème",
             'w:asciiTheme="majorHAnsi"' not in styles_xml.split('w:styleId="Heading1"')[1].split("</w:style>")[0])

    print("3. Page de garde et sommaire")
    verifier("un document long commence par une page de garde (titre, filet, date)",
             d.paragraphs[1].style.name == "Title" and "<w:pBdr>" in xml and any(m in xml for m in rendu._MOIS))
    verifier("le filet du titre est À SA PLACE (juste après le style, avant l'alignement)",
             '<w:pStyle w:val="Title"/><w:pBdr>' in xml and xml.index("<w:pBdr>") < xml.index('<w:jc', xml.index('<w:pStyle w:val="Title"/>')))
    verifier("en-tête et pied ne sont pas sur la page de garde", d.sections[0].different_first_page_header_footer)
    verifier("un sommaire (champ TOC) suit, dès cinq titres", 'TOC \\o "1-2"' in xml)
    verifier("Word met les champs à jour à l'ouverture, réglage À SA PLACE (avant compat)",
             "<w:updateFields" in reglages and reglages.index("<w:updateFields") < reglages.index("<w:compat"))

    print("4. Le texte et les tableaux")
    listes = [x for x in d.paragraphs if x.style.name == "List Bullet"]
    verifier("un paragraphe à tirets devient une vraie liste", len(listes) == 2 and listes[1].text == "Attestation d'assurance",
             [x.text for x in listes])
    verifier("« **gras** » devient du gras", any(r.bold and r.text == "Lettre de candidature" for r in listes[0].runs))
    table = d.tables[0]
    verifier("l'en-tête du tableau est à la couleur de la charte et se répète à chaque page",
             f'w:fill="{modele.COULEURS["charte"]}"' in xml and "<w:tblHeader" in xml)
    section = d.sections[0]
    utile = section.page_width - section.left_margin - section.right_margin
    largeurs = [int(g.get(qn("w:w"))) for g in table._tbl.tblGrid.findall(qn("w:gridCol"))]
    verifier("le tableau prend toute la largeur utile", abs(sum(largeurs) * 635 - utile) < utile * 0.02,
             (sum(largeurs) * 635, utile))

    print("5. Un document court reste simple")
    court = Document(rendre({"titre": "Note"}, [{"bloc": "paragraphe", "texte": "Bonjour."}]))
    verifier("ni page de garde ni sommaire pour une note",
             court.paragraphs[0].style.name == "Title" and "TOC" not in "".join(p.text for p in court.paragraphs)
             and not court.sections[0].different_first_page_header_footer)
    sans = Document(rendre({"titre": "Mémoire", "page_de_garde": False}, long))
    verifier("`page_de_garde: false` la retire même sur un long document",
             sans.paragraphs[0].style.name == "Title" and not sans.sections[0].different_first_page_header_footer)

    try:
        import reportlab  # noqa: F401
        pdf = tempfile.mktemp(suffix=".pdf")
        rendu._pdf(modele.normaliser_entete({"titre": "X", "format": "pdf"}),
                   [modele.normaliser_element({"bloc": "paragraphe", "texte": "a\n- b\n- c"})], pdf)
        verifier("le PDF rend un paragraphe à retours à la ligne", os.path.getsize(pdf) > 500)
    except ImportError:
        print("  (reportlab absent : rendu PDF sauté)")

src = (racine / "bureautique" / "rendu.py").read_text(encoding="utf-8")
verifier("le rendu Word part des styles de la maison", "_styles_maison(doc" in src and "_page_de_garde(doc" in src)

print()
if echecs:
    print(f"✗ {len(echecs)} échec(s)")
    sys.exit(1)
print("✓ 0 échec")
